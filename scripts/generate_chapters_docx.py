from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
DEFAULT_OUTPUT_PATH = ROOT / "CMU_ARLO_Chapters_1_to_3.docx"


TITLE = "A Web-Based Tracer and Employment Analytics System for the Alumni Relations and Linkages Office"
SHORT_TITLE = "CMU ARLO"


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def set_document_defaults(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def add_title_page(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Draft Chapters 1, 2, and 3")
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Prepared from the current CMU ARLO project documentation and system design artifacts")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    for _ in range(4):
        document.add_paragraph("")

    info = [
        "Institution: Central Mindanao University",
        "System Short Name: CMU ARLO",
        "Coverage: Chapter 1, Chapter 2, and Chapter 3",
        "Included Diagrams: ERD, IPO, and Data Flow Diagram Level 0",
    ]
    for line in info:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    document.add_page_break()


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles[f"Heading {level}"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    if level == 1:
        run.font.size = Pt(15)
        run.bold = True
    elif level == 2:
        run.font.size = Pt(13)
        run.bold = True
    else:
        run.font.size = Pt(12)
        run.bold = True


def add_paragraph(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text.strip())
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_figure(document: Document, image_path: Path, caption: str, width_inches: float = 6.0) -> None:
    if image_path.exists():
        document.add_picture(str(image_path), width=Inches(width_inches))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_paragraph(document, f"[Missing figure: {image_path.name}]")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


def add_ipo_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Input", "Process", "Output"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, "D9EAD3")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

    rows = [
        (
            "Alumni profile data, registration details, tracer responses, employment information, job applications, feedback entries, and administrative inputs such as announcements and job posting data.",
            "The system validates accounts, stores alumni records, processes tracer submissions, manages job applications, publishes announcements, routes notifications, and generates analytical summaries and audit records.",
            "Updated alumni records, portal views, application status results, published announcements, personalized notifications, management dashboards, analytical reports, and audit trails.",
        ),
        (
            "Login credentials, password reset requests, record maintenance requests, report filters, and settings changes from both alumni and administrators.",
            "Authentication, authorization, role-based access control, data retrieval, CRUD operations, aggregation of records, and presentation of summarized information for decision-making.",
            "Secure account access, verified registration results, profile updates, filtered reports, exportable analytics, and decision-support information for the Alumni Relations and Linkages Office.",
        ),
    ]

    for input_text, process_text, output_text in rows:
        row_cells = table.add_row().cells
        for idx, value in enumerate([input_text, process_text, output_text]):
            row_cells[idx].text = value
            for paragraph in row_cells[idx].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(11)


def build_document() -> Document:
    document = Document()
    set_document_defaults(document)
    add_title_page(document)

    # Chapter 1
    add_heading(document, "CHAPTER 1", 1)
    add_heading(document, "THE PROBLEM AND ITS BACKGROUND", 1)

    add_heading(document, "Introduction", 2)
    add_paragraph(
        document,
        "Alumni information is one of the most valuable institutional resources for higher education institutions because "
        "it provides evidence of graduate employability, supports curriculum improvement, and strengthens long-term "
        "engagement between the university and its former students. In many offices, however, alumni tracing is still "
        "performed through disconnected spreadsheets, social media follow-ups, paper forms, and manual consolidation of "
        "records. These approaches make reporting slow, inconsistent, and difficult to sustain over time."
    )
    add_paragraph(
        document,
        f"The {TITLE} is proposed to address these concerns by providing a centralized, web-based platform where alumni "
        "can register, maintain their profiles, submit tracer information, explore job opportunities, and send feedback, "
        "while administrators can manage records, publish announcements, review submissions, and generate analytics. "
        "Through this structure, the system supports the operational needs of the Alumni Relations and Linkages Office "
        "while also improving data accuracy, accessibility, and timeliness."
    )

    add_heading(document, "Background of the Study", 2)
    add_paragraph(
        document,
        "Graduate tracer systems are designed to monitor the educational and employment outcomes of alumni after they leave "
        "the university. They help institutions determine how well academic programs prepare graduates for employment, what "
        "industries absorb graduates, and what improvements may be needed in curriculum, services, and extension programs. "
        "For the Alumni Relations and Linkages Office, the absence of a centralized digital system can create repeated data "
        "entry, scattered records, limited reporting capability, and delays in responding to alumni concerns."
    )
    add_paragraph(
        document,
        "The current CMU ARLO project already demonstrates the feasibility of such a platform. Based on the project "
        "documentation and codebase, the system includes authentication, alumni profile management, tracer data handling, "
        "job board workflows, feedback channels, announcements, notifications, analytics, and audit logging. These modules "
        "reflect the core functions needed by the office to manage alumni records and produce timely employment insights."
    )

    add_heading(document, "Statement of the Problem", 2)
    add_paragraph(
        document,
        "This study is guided by the general problem of how to design and implement a web-based tracer and employment analytics "
        "system that will improve the efficiency, accuracy, and usefulness of alumni records management for the Alumni Relations "
        "and Linkages Office."
    )
    add_paragraph(document, "Specifically, it seeks to answer the following questions:")
    add_numbered(
        document,
        [
            "How can alumni registration, profile maintenance, and tracer data submission be centralized in a single web-based platform?",
            "How can the system support job posting, job application, feedback submission, and alumni engagement activities?",
            "How can administrators efficiently manage announcements, submissions, records, and reports using role-based access controls?",
            "How can the system generate useful employment and tracer analytics that can support planning and decision-making?",
            "How can the proposed database and process design maintain data consistency, privacy, and traceability?"
        ],
    )

    add_heading(document, "Objectives of the Study", 2)
    add_paragraph(
        document,
        "The general objective of the study is to develop a web-based tracer and employment analytics system for the Alumni "
        "Relations and Linkages Office."
    )
    add_paragraph(document, "Specifically, the study aims to:")
    add_bullets(
        document,
        [
            "Provide a secure account registration and authentication mechanism for alumni and administrators.",
            "Allow alumni to maintain personal and employment records through a structured tracer workflow.",
            "Provide modules for jobs, announcements, feedback, and notifications to strengthen alumni engagement.",
            "Enable administrators to manage alumni records, submissions, announcements, and reports from a central dashboard.",
            "Produce analytical outputs that summarize tracer and employment information for institutional use.",
            "Design a database structure and process flow that support consistent and auditable information management.",
        ],
    )

    add_heading(document, "Significance of the Study", 2)
    add_paragraph(
        document,
        "The study is significant to several stakeholders. For the Alumni Relations and Linkages Office, it offers a more "
        "organized and efficient way of collecting, maintaining, and analyzing alumni information. For university administrators, "
        "it can support data-driven planning, policy formulation, and reporting. For alumni, it provides a convenient channel "
        "for registration, profile updating, job exploration, and communication with the university. For future researchers and "
        "developers, the study can serve as a reference for improving alumni information systems, tracer systems, and "
        "employment analytics platforms."
    )

    add_heading(document, "Scope and Delimitation", 2)
    add_paragraph(
        document,
        f"The scope of the study covers the design and documentation of the {SHORT_TITLE} platform for the Alumni Relations "
        "and Linkages Office. The system includes two major user roles, namely administrator and alumni. Its main functional "
        "coverage includes registration, authentication, profile and settings management, tracer record submission, job browsing "
        "and application, job opportunity submission, announcements, notifications, feedback handling, analytics, and audit logging."
    )
    add_paragraph(
        document,
        "The study is limited to a web-based implementation and the data structures observed in the current repository. It does "
        "not include external enterprise integrations such as payroll systems, third-party human resource platforms, or fully "
        "automated labor market data feeds. The content of the chapter draft is also based on the present project documents and "
        "system modules rather than externally cited literature."
    )

    add_heading(document, "Definition of Terms", 2)
    add_bullets(
        document,
        [
            "Alumni Tracer System: A system used to collect and analyze information about graduates, particularly their employment status, professional outcomes, and educational experiences after graduation.",
            "Employment Analytics: The process of summarizing and interpreting employment-related data to support planning, evaluation, and decision-making.",
            "Administrator: A user role responsible for managing records, announcements, job postings, reports, and system oversight functions.",
            "Alumni: A graduate or former student who accesses the system to register, manage profile information, submit tracer data, and participate in alumni services.",
            "Notification: A system-generated message used to inform users about updates such as announcements, approvals, or reminders.",
            "Audit Log: A record of important system activities used for monitoring, accountability, and traceability.",
        ],
    )

    document.add_page_break()

    # Chapter 2
    add_heading(document, "CHAPTER 2", 1)
    add_heading(document, "REVIEW OF RELATED CONCEPTS AND SYSTEM FRAMEWORK", 1)

    add_heading(document, "Overview of Related Concepts", 2)
    add_paragraph(
        document,
        "A web-based alumni tracer platform combines several important information system concepts. These include centralized "
        "data management, user authentication, role-based access control, record processing, reporting, and feedback management. "
        "In the context of higher education, a tracer system becomes more valuable when it does not only store alumni records "
        "but also transforms them into usable information for planning and evaluation."
    )
    add_paragraph(
        document,
        "The CMU ARLO project reflects this direction by combining alumni engagement functions with analytical and administrative "
        "tools. Instead of treating alumni tracing as an isolated survey activity, the system positions it as part of a broader "
        "digital environment that includes announcements, job opportunities, notifications, and office monitoring functions."
    )

    add_heading(document, "Alumni Tracer Systems", 2)
    add_paragraph(
        document,
        "Alumni tracer systems are designed to monitor graduates over time. Their essential role is to gather data on graduate "
        "employment, continuing education, relevance of curriculum, and professional experiences. Such systems are important "
        "because they help identify whether a university program remains aligned with labor market demands. In practice, a good "
        "tracer system must make it easy for alumni to respond and equally easy for administrators to organize and interpret the results."
    )

    add_heading(document, "Employment Analytics", 2)
    add_paragraph(
        document,
        "Employment analytics refers to the collection, organization, and interpretation of employment-related data for reporting "
        "and decision-making. Within an alumni office, this may include employment status, occupation, industry, salary range, "
        "job relevance to degree programs, and the time it takes graduates to obtain their first job. By transforming raw tracer "
        "submissions into summarized reports, the institution gains evidence that can be used in accreditation, program review, "
        "and strategic planning."
    )

    add_heading(document, "Web-Based Information Systems", 2)
    add_paragraph(
        document,
        "Web-based information systems provide accessibility, centralized storage, and easier maintenance compared with fragmented "
        "manual processes. Because users can access them through standard browsers, they reduce the need for device-specific "
        "installations and allow both alumni and administrators to interact with shared records in real time. In the present "
        "project, the web-based approach supports multiple modules from a single interface, including authentication, tracer "
        "submission, job management, and analytics."
    )

    add_heading(document, "Role-Based Access and Data Privacy", 2)
    add_paragraph(
        document,
        "Role-based access control is essential in systems where different user groups have different responsibilities. Alumni "
        "need access to their own data, announcements, and employment-related functions, while administrators need access to "
        "management, reporting, and audit functions. Data privacy is also a crucial consideration because the system stores "
        "personally identifiable information and employment-related records. The current project includes consent-aware registration "
        "and settings controls, which support privacy-conscious system use."
    )

    add_heading(document, "Conceptual Framework", 2)
    add_paragraph(
        document,
        "The conceptual framework of the proposed study can be presented using the Input-Process-Output (IPO) model. The IPO model "
        "explains how user-provided data and administrative inputs are transformed by the system into structured records, services, "
        "and analytical outputs. This framework is appropriate because the project revolves around data collection, processing, "
        "storage, and reporting."
    )

    add_heading(document, "Input-Process-Output (IPO) Framework", 2)
    add_ipo_table(document)
    add_paragraph(
        document,
        "The input portion of the framework consists of all information submitted to the system by alumni and administrators. "
        "These inputs include registration details, tracer responses, employment information, feedback, job submissions, and "
        "administrative maintenance activities. The process portion represents the major internal system functions such as account "
        "verification, authentication, CRUD operations, data validation, analytics generation, and notification routing. The "
        "output portion includes the visible results of processing, such as updated records, dashboards, job application statuses, "
        "announcements, and reports. In this way, the IPO framework demonstrates how the proposed system converts raw alumni and "
        "administrative inputs into meaningful and usable outputs for decision-making."
    )
    add_paragraph(
        document,
        "In a more detailed sense, the input component of the framework includes the personal, academic, and employment data "
        "provided by alumni during registration and tracer submission. It also includes administrative inputs such as announcement "
        "content, job posting details, review decisions, report filters, and maintenance updates. These inputs are important "
        "because they serve as the raw data that the system must capture, organize, and protect."
    )
    add_paragraph(
        document,
        "The process component represents the internal activities of the system. These include verifying registration details, "
        "authenticating users, validating submitted forms, saving and updating records, reviewing job-related submissions, "
        "publishing announcements, routing notifications, summarizing tracer data, and logging critical activities for audit "
        "purposes. Through these processes, the system converts unstructured user inputs into organized and reliable information."
    )
    add_paragraph(
        document,
        "The output component refers to the results that become available to users after processing takes place. For alumni, "
        "these outputs may include confirmed account access, updated profile records, available job opportunities, application "
        "status updates, notifications, and tracer submission confirmations. For administrators, outputs include managed alumni "
        "records, published announcements, feedback monitoring views, audit trails, and analytical reports that summarize alumni "
        "employment and tracer information. These outputs help the office carry out operations more efficiently while also "
        "supporting institutional decision-making."
    )
    add_heading(document, "Detailed Explanation of the IPO Framework", 2)
    add_paragraph(
        document,
        "The IPO framework is highly appropriate for the proposed system because the platform is fundamentally concerned with "
        "collecting data, processing it through defined system functions, and presenting meaningful results to different users. "
        "It shows that the system does more than merely store records. Instead, it transforms alumni and administrative data into "
        "practical outputs such as dashboards, reports, notifications, and decision-support information. For this reason, the IPO "
        "model provides a clear conceptual guide for understanding the behavior and purpose of the proposed system."
    )

    add_heading(document, "Chapter Summary", 2)
    add_paragraph(
        document,
        "Chapter 2 discussed the concepts that support the development of the proposed web-based tracer and employment analytics "
        "system. It highlighted the role of tracer systems, employment analytics, web-based platforms, role-based access control, "
        "and data privacy. It also presented the Input-Process-Output framework as the conceptual guide for understanding how the "
        "system transforms data into services and reports."
    )

    document.add_page_break()

    # Chapter 3
    add_heading(document, "CHAPTER 3", 1)
    add_heading(document, "METHODOLOGY AND SYSTEM DESIGN", 1)

    add_heading(document, "Research and Development Approach", 2)
    add_paragraph(
        document,
        "This study uses the Waterfall model as the software development methodology for the proposed system. The Waterfall model "
        "is a linear and sequential approach in which each phase of development is completed before the next phase begins. It is "
        "appropriate for the study because the major requirements of the Alumni Relations and Linkages Office can be identified "
        "early, documented clearly, and translated into formal design outputs such as process models, database structures, and "
        "system modules."
    )
    add_paragraph(
        document,
        "The Waterfall model is also appropriate in the context of a thesis or capstone study because it supports thorough "
        "documentation at every stage. Since the proposed system requires a formal presentation of requirements, diagrams, design "
        "structures, implementation details, and verification activities, the Waterfall model provides a clear sequence that is "
        "easy to explain, monitor, and justify."
    )

    add_heading(document, "Phases of the Waterfall Model Used in the Study", 2)
    add_numbered(
        document,
        [
            "Requirements Analysis. In this phase, the system requirements of the Alumni Relations and Linkages Office were identified. These include alumni registration, profile management, tracer data collection, job management, announcements, notifications, feedback handling, and analytics.",
            "System Design. After identifying the requirements, the structure of the system was designed. This phase includes the preparation of the database design, entity relationship diagram, data flow diagrams, and the identification of modules and user roles.",
            "Implementation. In this phase, the planned design was translated into an actual web-based platform using Laravel, Blade templates, PostgreSQL, Vite, Tailwind CSS, and related project tools.",
            "Testing. After implementation, the major functions of the system were checked to ensure correctness. Based on the current repository, testing covers important workflows such as authentication, feedback handling, CRUD operations, job applications, and account-related actions.",
            "Deployment. Once the system is stable and validated, it may be prepared for actual use in the office environment through local or server-based deployment.",
            "Maintenance. After deployment, the system can be improved through updates, revisions, bug fixes, and additional features based on user needs and organizational changes."
        ],
    )
    add_paragraph(
        document,
        "Through the Waterfall model, the development of the proposed system follows a structured path from requirement gathering "
        "up to testing and maintenance. This makes the methodology suitable for a formal academic study where every stage must be "
        "clearly presented and supported by documentation."
    )

    add_heading(document, "System Architecture Overview", 2)
    add_paragraph(
        document,
        "Based on the current project repository, the system uses Laravel as the main application framework, Blade templates for "
        "the user interface, PostgreSQL for live application data, SQLite in memory for automated testing, Vite for asset handling, "
        "and supporting libraries such as Chart.js and Boxicons. This architecture supports both transactional operations and "
        "presentation-oriented modules needed for an alumni portal and administrative dashboard."
    )

    add_heading(document, "Major Functional Modules", 2)
    add_bullets(
        document,
        [
            "Authentication and account management",
            "Alumni profile and settings management",
            "Tracer and employment data collection",
            "Jobs, applications, and job submission workflows",
            "Announcements and notifications",
            "Feedback monitoring and response",
            "Analytics, reports, and audit logging",
        ],
    )

    add_heading(document, "Data Flow Diagram Level 0", 2)
    add_paragraph(
        document,
        "The Data Flow Diagram (DFD) Level 0, or context diagram, presents the proposed system as a single major process and "
        "shows how it interacts with its two main external entities: the alumni and the administrator."
    )
    add_figure(
        document,
        DOCS_DIR / "arlo-dfd-context-hd.png",
        "Figure 1. Data Flow Diagram Level 0 of the Proposed Web-Based Tracer and Employment Analytics System",
        width_inches=6.2,
    )
    add_paragraph(
        document,
        "In the diagram, alumni provide inputs such as login credentials, profile updates, tracer submissions, job applications, "
        "and feedback. In return, they receive outputs such as portal content, announcements, notifications, jobs, and their "
        "personal record information. On the other hand, the administrator provides management inputs such as record maintenance, "
        "job posting, announcement publishing, and analytics requests. The system then returns dashboards, managed records, "
        "reports, and activity logs. The Level 0 DFD therefore shows the overall communication boundary of the system without yet "
        "breaking down its internal processes."
    )

    add_heading(document, "Entity Relationship Diagram", 2)
    add_paragraph(
        document,
        "The Entity Relationship Diagram (ERD) illustrates the structure of the database and the relationships among the main "
        "entities used by the system."
    )
    add_figure(
        document,
        DOCS_DIR / "arlo-erd-hd.png",
        "Figure 2. Entity Relationship Diagram of the Proposed Web-Based Tracer and Employment Analytics System",
        width_inches=6.2,
    )
    add_paragraph(
        document,
        "The ERD shows that the Users entity is central to the database because it connects to tracer records, job applications, "
        "job submissions, announcements, feedback, notifications, system logs, and sessions. Tracer records store the employment "
        "and graduate outcome details linked to a specific alumnus. Jobs and job applications support recruitment and alumni "
        "engagement features, while announcements and notifications support information dissemination. Feedback, sessions, and "
        "system logs provide monitoring and administrative support functions. This structure helps ensure that the system can "
        "store both operational records and analytical data in a consistent manner."
    )

    add_heading(document, "Database Design Explanation", 2)
    add_paragraph(
        document,
        "The database design follows a relational structure in which each major module of the system is represented by a table or "
        "entity. The user table stores core alumni and administrator information. Tracer records are linked to users to allow one "
        "alumnus to have one or more tracer entries over time. Job applications link alumni to job records, while announcements "
        "and feedback create separate but related channels for communication and office services. This design supports data "
        "integrity, minimizes duplication, and makes it easier to generate filtered summaries for analytics and reporting."
    )

    add_heading(document, "Processing and Validation Strategy", 2)
    add_paragraph(
        document,
        "The proposed system uses validation and controlled workflows to improve record quality. Registration includes verification "
        "logic, account credentials are processed securely, and role-based functions limit the operations available to each user "
        "type. Tracer submissions, feedback entries, and job-related transactions are processed through forms and status-based "
        "operations. These strategies help ensure that the information saved in the system remains reliable and usable for future reporting."
    )

    add_heading(document, "Testing and Quality Considerations", 2)
    add_paragraph(
        document,
        "The repository includes feature tests for authentication, admin CRUD flows, audit logs, feedback, alumni job applications, "
        "announcement visibility, and account deletion cascade behavior. This indicates that the project values verification of "
        "critical workflows. In a thesis context, these tests support the claim that the system design is not only documented but "
        "also checked for correctness across important user actions."
    )

    add_heading(document, "Chapter Summary", 2)
    add_paragraph(
        document,
        "Chapter 3 presented the methodology and system design for the proposed platform. It described the development approach, "
        "outlined the major modules, explained the Level 0 DFD and ERD, and discussed the database, processing, and testing "
        "considerations of the system. Together, these elements provide a technical foundation for the implementation of the "
        "web-based tracer and employment analytics system."
    )

    return document


def main() -> None:
    document = build_document()
    output_path = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_OUTPUT_PATH
    document.save(str(output_path))
    print(f"Created: {output_path}")


if __name__ == "__main__":
    main()
