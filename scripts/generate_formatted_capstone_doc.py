from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
REFERENCE_PATH = Path(
    r"c:\Users\washe\OneDrive\washy\ARLO(CAPSTONE)\A Web-Based Tracer and Employment Analytics System for Alumni Relation and Linkages Office.docx"
)
DEFAULT_OUTPUT_PATH = REFERENCE_PATH.with_name(
    "A Web-Based Tracer and Employment Analytics System for Alumni Relation and Linkages Office_Revised_Formatted.docx"
)

REVIEWED_WORKS = [
    {
        "label": "Centralized Alumni System (Kumar & Thomas, 2019)",
        "feature": "Centralized alumni database and searchable records",
        "contribution": "Inspired the centralized alumni record and profile management approach of the proposed system.",
    },
    {
        "label": "Secure Alumni Portal (Goyal & Goyal, 2017)",
        "feature": "Protected login and controlled user access",
        "contribution": "Guided the use of secure authentication and role-based access for alumni and administrators.",
    },
    {
        "label": "Cloud-Based Secure Alumni Portal (Saini & Bala, 2018)",
        "feature": "Online alumni services with strong data protection",
        "contribution": "Reinforced privacy, controlled access, and secure record handling in a web environment.",
    },
    {
        "label": "Smart Alumni Portal (Chintalapati & Pandharipande, 2020)",
        "feature": "Web communication and alumni engagement tools",
        "contribution": "Influenced the announcements, notifications, and alumni engagement modules of the study.",
    },
]

REFERENCES = [
    "Chintalapati, S., & Pandharipande, R. (2020). Smart Alumni Portal for Higher Educational Institutions Using Cloud Computing. International Journal of Computer Science and Mobile Computing, 9(5), 38-45.",
    "Goyal, R., & Goyal, R. (2017). Design and Implementation of Secure Alumni Portal. International Journal of Computer Applications, 161(8), 25-29.",
    "Kumar, R., & Thomas, M. (2019). Design and Development of Alumni Management System for Higher Education Institutions. International Journal of Scientific Research in Computer Science, Engineering and Information Technology, 5(2), 225-230.",
    "Saini, H., & Bala, A. (2018). Design and Implementation of a Cloud-Based Secure Alumni Portal. International Journal of Computer Sciences and Engineering, 6(3), 67-72.",
]


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def clear_document(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def configure_document(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(12)

    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def parse_reference_front_matter(reference: Document) -> dict[str, object]:
    centered_lines: list[str] = []
    for paragraph in reference.paragraphs:
        text = paragraph.text.strip()
        if "CHAPTER I" in text:
            break
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER and text:
            centered_lines.append(text)

    title = centered_lines[0] if centered_lines else "A Web-Based Tracer and Employment Analytics System for Alumni Relation and Linkages Office"
    authors: list[str] = []
    course = "Bachelor Of Science in Information Technology"
    date_text = "MARCH 2026"

    for line in centered_lines[1:]:
        if line.upper().startswith("CHAPTER"):
            continue
        if "Bachelor" in line or "BS" in line:
            course = line
            continue
        if line.upper() == line and any(char.isdigit() for char in line):
            date_text = line
            continue
        authors.append(line)

    if not authors:
        authors = ["Washeley Tan", "Elvin Lar Gonzales", "Jose Carlo Benegian"]

    return {
        "title": title,
        "authors": authors,
        "course": course,
        "date_text": date_text,
    }


def add_paragraph(
    document: Document,
    text: str = "",
    *,
    alignment: WD_ALIGN_PARAGRAPH | None = WD_ALIGN_PARAGRAPH.JUSTIFY,
    bold: bool = False,
    italic: bool = False,
    font_size: float = 12,
    line_spacing: float = 1.5,
) -> None:
    paragraph = document.add_paragraph()
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = line_spacing
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic


def add_empty_line(document: Document, *, alignment: WD_ALIGN_PARAGRAPH | None = WD_ALIGN_PARAGRAPH.CENTER) -> None:
    paragraph = document.add_paragraph()
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.5


def add_title_page(document: Document, meta: dict[str, object]) -> None:
    for _ in range(5):
        add_empty_line(document)

    add_paragraph(
        document,
        str(meta["title"]),
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
    )

    for _ in range(5):
        add_empty_line(document)

    for author in meta["authors"]:
        add_paragraph(document, str(author), alignment=WD_ALIGN_PARAGRAPH.CENTER)

    for _ in range(3):
        add_empty_line(document)

    add_paragraph(document, str(meta["course"]), alignment=WD_ALIGN_PARAGRAPH.CENTER)

    for _ in range(5):
        add_empty_line(document)

    add_paragraph(document, str(meta["date_text"]), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    document.add_page_break()


def add_chapter_heading(document: Document, chapter: str, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(f"{chapter}\n{title}")
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.bold = True


def add_section_heading(document: Document, text: str) -> None:
    add_paragraph(document, text, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=True)


def add_numbered_items(document: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        add_paragraph(document, f"{index}. {item}")


def add_hanging_items(document: Document, items: list[str], prefix: str = "- ") -> None:
    for item in items:
        add_paragraph(document, f"{prefix}{item}")


def add_review_table(document: Document) -> None:
    add_paragraph(document, "Table 2.1 Summary of Reviewed Systems and Studies", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["System/Study", "Key Feature", "Contribution to Proposed System"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, "D9EAD3")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(11)
                run.bold = True

    for work in REVIEWED_WORKS:
        row = table.add_row().cells
        values = [work["label"], work["feature"], work["contribution"]]
        for index, value in enumerate(values):
            row[index].text = value
            for paragraph in row[index].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(11)


def add_ipo_table(document: Document) -> None:
    add_paragraph(document, "Table 2.2 Input-Process-Output Framework", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Input", "Process", "Output"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, "D9EAD3")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(11)
                run.bold = True

    rows = [
        (
            "Alumni personal, academic, and employment information; registration details; tracer responses; job applications; feedback; and administrator maintenance inputs.",
            "Account verification, form validation, storage and updating of records, review of submissions, notification routing, report generation, and audit logging.",
            "Verified accounts, updated alumni records, tracer summaries, published announcements, application status updates, dashboards, and office reports.",
        ),
        (
            "Job postings, announcement content, report filters, settings changes, and communication requests from administrators and alumni.",
            "Publishing, audience filtering, status monitoring, role-based access control, aggregation of data, and presentation of summaries for decision-making.",
            "Relevant announcements, targeted notifications, employment charts, filtered reports, and decision-support information for the Alumni Relation and Linkages Office.",
        ),
    ]

    for input_text, process_text, output_text in rows:
        row = table.add_row().cells
        for index, value in enumerate([input_text, process_text, output_text]):
            row[index].text = value
            for paragraph in row[index].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(11)


def add_figure(document: Document, image_path: Path, caption: str, *, width_inches: float = 6.0) -> None:
    if image_path.exists():
        document.add_picture(str(image_path), width=Inches(width_inches))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_paragraph(document, f"[Missing figure: {image_path.name}]")

    add_paragraph(
        document,
        caption,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        italic=True,
        font_size=11,
    )


def build_document() -> Document:
    reference = Document(str(REFERENCE_PATH))
    meta = parse_reference_front_matter(reference)

    document = Document()
    configure_document(document)
    add_title_page(document, meta)

    add_chapter_heading(document, "CHAPTER I", "INTRODUCTION")
    add_section_heading(document, "1. Background of the Study")
    add_paragraph(
        document,
        "Maintaining accurate alumni records is important for every university because graduate information supports communication, tracer activities, employment monitoring, and institutional planning. When alumni records are scattered across paper forms, spreadsheets, social media messages, and isolated files, the office spends more time consolidating information than analyzing it. This weakens the ability of the institution to respond quickly to alumni concerns and to produce timely reports about graduate outcomes."
    )
    add_paragraph(
        document,
        "The Alumni Relation and Linkages Office (ARLO) of Central Mindanao University needs a more organized and responsive way to manage alumni information. Aside from keeping graduate profiles, the office also needs to monitor employment status, publish opportunities, share announcements, and collect feedback from alumni. Manual and fragmented processes make these tasks difficult because records become inconsistent, updates are delayed, and reports require additional effort to prepare."
    )
    add_paragraph(
        document,
        "To address these concerns, this study proposes a web-based tracer and employment analytics system for ARLO. The system centralizes alumni records, employment-related submissions, job opportunities, announcements, notifications, and feedback in one platform. Through a centralized and role-aware environment, both alumni and administrators can interact with the same institutional data more efficiently."
    )
    add_paragraph(
        document,
        "The proposed system also emphasizes analytics and reporting. Instead of serving only as a storage platform, it helps ARLO visualize graduate outcomes through employment summaries, filtered reports, and dashboard-based indicators. These outputs can support office operations, strategic planning, and future curriculum or alumni-service improvements."
    )

    add_section_heading(document, "1.2 Statement of the Problem")
    add_paragraph(
        document,
        "The Alumni Relation and Linkages Office (ARLO) of Central Mindanao University needs an efficient and organized way to manage alumni information, monitor employment outcomes, and communicate important updates to graduates. Traditional and fragmented record-keeping makes it difficult to maintain accurate data, generate timely reports, and support tracer-related services."
    )
    add_paragraph(document, "Specifically, the study seeks to answer the following questions:")
    add_numbered_items(
        document,
        [
            "How can alumni registration, profile maintenance, and tracer data submission be centralized in one web-based platform?",
            "How can the system support jobs, feedback, announcements, and other engagement functions needed by alumni and ARLO personnel?",
            "How can administrators manage records, submissions, and communication efficiently through a role-based interface?",
            "How can the system generate useful employment and tracer analytics for reporting and decision-making?",
            "How can the proposed system maintain data consistency, privacy, and traceability during day-to-day operations?",
        ],
    )

    add_section_heading(document, "1.3 Objectives of the Study")
    add_paragraph(
        document,
        "The general objective of the study is to develop a web-based tracer and employment analytics system for the Alumni Relation and Linkages Office that centralizes records, improves communication, and supports employment-related reporting."
    )
    add_paragraph(document, "Specific objectives:")
    add_hanging_items(
        document,
        [
            "To develop a centralized web-based system for managing alumni personal, academic, and employment records.",
            "To enable alumni to update their profile and tracer-related information through an online portal.",
            "To provide administrators with tools for managing alumni records, job postings, announcements, notifications, and feedback.",
            "To generate employment and tracer analytics through dashboards, charts, and filtered reports.",
            "To support targeted communication by allowing announcements and notifications to reach the appropriate alumni groups.",
            "To implement secure role-based access and protected handling of alumni data.",
        ],
    )

    add_section_heading(document, "1.4 Scope and Limitation of the Study")
    add_paragraph(
        document,
        "This study focuses on the design and development of a web-based tracer and employment analytics system for the Alumni Relation and Linkages Office of Central Mindanao University. The system allows alumni to register, maintain verified profile information, submit tracer details, access jobs and announcements, receive notifications, and send feedback. On the administrative side, the platform supports alumni record management, review workflows, communication tools, and analytics dashboards."
    )
    add_paragraph(
        document,
        "The study is limited to a web-based implementation and does not include a standalone mobile application. It is designed specifically for the workflow and data requirements observed in the current CMU ARLO project. External integrations with third-party recruitment platforms, SMS gateways, or university-wide enterprise systems are outside the present scope. The study also focuses on the features available in the current repository and related documentation."
    )

    add_section_heading(document, "1.5 Significance of the Study")
    add_paragraph(
        document,
        "This study is significant because it provides a practical and organized solution for alumni tracing, employment monitoring, and communication management in the university setting."
    )
    add_paragraph(
        document,
        "For the Alumni Relation and Linkages Office (ARLO): The system helps reduce manual work in maintaining and retrieving alumni records. It provides a centralized database, supports targeted communication, and offers employment analytics that can improve reporting and planning."
    )
    add_paragraph(
        document,
        "For CMU Alumni: The system offers a more convenient way to maintain records, receive announcements, access job opportunities, submit feedback, and participate in tracer-related activities through one online portal."
    )
    add_paragraph(
        document,
        "For CMU Administration and Future Researchers: The system can provide clearer insights into graduate outcomes, employment trends, and alumni engagement. The study may also serve as a reference for future work involving alumni information systems, tracer studies, employment analytics, and institutional web platforms."
    )

    document.add_page_break()

    add_chapter_heading(document, "CHAPTER II", "REVIEW OF RELATED LITERATURE")
    add_section_heading(document, "2.1 Review of Related Systems")
    add_paragraph(
        document,
        "To guide the development of the proposed system, existing alumni management and portal-based systems were reviewed. These systems show how centralized data handling, secure access, and alumni self-service can improve the management of graduate records."
    )
    add_paragraph(
        document,
        "Kumar and Thomas (2019) presented an alumni management system that emphasized a centralized database for storing and retrieving graduate information. Their work highlighted the value of searchable alumni records and organized data maintenance, which are also important in the proposed study."
    )
    add_paragraph(
        document,
        "Goyal and Goyal (2017) focused on the design of a secure alumni portal with attention to login mechanisms, controlled access, and data protection. Their study is relevant because the present system also requires secure authentication and differentiated access for alumni and administrators."
    )
    add_paragraph(
        document,
        "Saini and Bala (2018) discussed a cloud-based secure alumni portal and reinforced the importance of privacy, protected transactions, and structured online services. Their work supports the need for secure and dependable handling of alumni information in a web-based environment."
    )
    add_paragraph(
        document,
        "Chintalapati and Pandharipande (2020) described a smart alumni portal that strengthened alumni engagement through web-based services and communication support. This aligns with the proposed system because ARLO also needs announcements, notifications, and engagement features alongside tracer-related functions."
    )

    add_section_heading(document, "2.2 Review of Related Studies")
    add_paragraph(
        document,
        "Related studies also show that alumni systems become more effective when users can update their own records and when administrators can generate meaningful reports from submitted data. In many institutions, the value of an alumni system does not only come from storing profiles but from converting alumni responses into insights that help the university evaluate outcomes and plan improvements."
    )
    add_paragraph(
        document,
        "The reviewed literature consistently shows three recurring needs. First, institutions need a centralized structure for alumni records (Kumar & Thomas, 2019). Second, they need secure and controlled access to protect user data (Goyal & Goyal, 2017; Saini & Bala, 2018). Third, they need web-based communication and engagement functions so that alumni can actively participate and remain informed (Chintalapati & Pandharipande, 2020)."
    )
    add_paragraph(
        document,
        "The proposed system extends these ideas by placing stronger emphasis on tracer and employment analytics. Instead of limiting the platform to profile storage, the study combines alumni registration, verified records, employment-related inputs, communication tools, and dashboards in one system. This integration helps ARLO not only maintain information but also interpret it for institutional use."
    )
    add_paragraph(
        document,
        "Most of the reviewed systems were discussed in broader or international contexts. Because university processes vary by institution, there is value in a system that reflects the needs and workflow of the Central Mindanao University Alumni Relation and Linkages Office. The present study therefore adapts common best practices to the local context."
    )

    add_section_heading(document, "2.3 Summary of Related Systems and Studies")
    add_paragraph(
        document,
        "The reviewed systems and studies guided the design of the proposed project by highlighting the importance of centralized records, secure access, self-service updates, and reporting functions. These findings support the direction of the present system and justify its combination of tracer management, communication, and analytics modules."
    )
    add_review_table(document)
    add_paragraph(
        document,
        "Overall, the literature review establishes a strong foundation for the proposed system. The reviewed systems show that alumni platforms are more effective when they are secure, centralized, and supportive of engagement and reporting. These lessons are directly reflected in the design of the proposed ARLO platform."
    )

    add_section_heading(document, "2.4 Conceptual Framework (IPO)")
    add_paragraph(
        document,
        "The proposed tracer and employment analytics system follows an Input-Process-Output (IPO) structure. The IPO framework is appropriate because the platform collects structured information from alumni and administrators, transforms that information through system functions, and produces records, services, and reports that support ARLO operations."
    )
    add_ipo_table(document)
    add_paragraph(
        document,
        "In the input stage, the system receives personal, academic, and employment details from alumni, along with administrative inputs such as job postings, announcements, report filters, and record maintenance actions. These inputs become the raw data needed for tracer management, communication, and employment monitoring."
    )
    add_paragraph(
        document,
        "In the process stage, the system validates registration details, authenticates users, stores or updates records, reviews submissions, routes notifications, and aggregates tracer information. This stage is where the platform transforms raw form entries into organized and reliable institutional data."
    )
    add_paragraph(
        document,
        "In the output stage, the system produces verified accounts, updated profiles, tracer summaries, dashboards, job application status updates, published announcements, notifications, and reports. These outputs help ARLO maintain accurate alumni information and make better decisions based on current data."
    )
    add_paragraph(
        document,
        "In summary, the system transforms raw alumni, employment, and administrative inputs into organized records and useful outputs. This structure allows ARLO to maintain better records, communicate more effectively, and generate more meaningful employment and tracer analytics."
    )

    document.add_page_break()

    add_chapter_heading(document, "CHAPTER III", "TECHNICAL BACKGROUND")
    add_paragraph(
        document,
        "The technical background of the study describes the technologies, system structure, and design artifacts that support the proposed platform. The project is implemented as a web-based system so that alumni and administrators can access the same centralized environment through a browser-based interface."
    )
    add_section_heading(document, "Tools Used")
    tools_used = [
        "Laravel (PHP Framework) - Laravel serves as the main application framework. It organizes routing, validation, controllers, database operations, and modular service logic for both alumni and administrator features.",
        "PHP 8.2 (Programming Language) - PHP powers the server-side logic of the application, including authentication, verification, submission processing, analytics preparation, and notification handling.",
        "PostgreSQL (Database) - PostgreSQL stores live system data such as users, tracer records, jobs, announcements, notifications, feedback, and audit-related records.",
        "SQLite (Testing Database) - SQLite is used for automated tests so that feature flows can be verified quickly without affecting the live database.",
        "Blade Templates, Tailwind CSS, and Vite - Blade builds the application views, Tailwind CSS supports layout and styling, and Vite handles front-end asset bundling.",
        "Chart.js - Chart.js is used to display employment trends, graduation distributions, tracer summaries, and analytics dashboards.",
        "Notification and Mail Support - The system supports in-app and email notifications for announcements, reviews, and account-related actions through Laravel notification services.",
    ]
    add_hanging_items(document, tools_used)

    add_section_heading(document, "System Architecture Overview")
    add_paragraph(
        document,
        "Based on the current project repository, the system uses Laravel as the main framework, Blade for the user interface, PostgreSQL for primary data storage, SQLite for automated testing, and front-end tooling that includes Tailwind CSS, Vite, and Chart.js. This architecture supports both operational workflows and presentation-oriented analytics needed by ARLO."
    )

    add_section_heading(document, "Major Functional Modules")
    add_hanging_items(
        document,
        [
            "Authentication and account management",
            "Alumni profile and settings management",
            "Tracer and employment data collection",
            "Jobs, applications, and job submission workflows",
            "Announcements and notifications",
            "Feedback monitoring and response",
            "Analytics, reports, and audit logging",
        ],
    )

    add_section_heading(document, "Entity Relationship Diagram (ERD)")
    add_paragraph(
        document,
        "The Entity Relationship Diagram illustrates the structure of the database and the relationships among the main records used by the system."
    )
    add_figure(
        document,
        DOCS_DIR / "arlo-erd-hd.png",
        "Figure 3. Entity Relationship Diagram of the Proposed Web-Based Tracer and Employment Analytics System",
        width_inches=6.2,
    )
    add_paragraph(
        document,
        "The ERD shows that the user record is central to the database because it connects to tracer entries, job applications, job submissions, announcements, notifications, feedback, sessions, and system logs. This structure helps the system maintain organized relationships among personal records, communication features, and employment-related transactions."
    )
    add_paragraph(
        document,
        "The relational design supports data integrity by reducing duplication and by linking records through defined relationships. This makes it easier to retrieve records accurately, monitor user activity, and generate filtered reports for tracer and employment analytics."
    )

    add_section_heading(document, "Data Flow Diagram Level 0")
    add_paragraph(
        document,
        "The Data Flow Diagram Level 0, also called the context diagram, presents the proposed system as one major process and shows how it interacts with the two main external entities: alumni and administrators."
    )
    add_figure(
        document,
        DOCS_DIR / "arlo-dfd-context-hd.png",
        "Figure 4. Data Flow Diagram Level 0 of the Proposed Web-Based Tracer and Employment Analytics System",
        width_inches=6.2,
    )
    add_paragraph(
        document,
        "In the diagram, alumni provide registration details, profile updates, tracer responses, job applications, and feedback to the system. In return, they receive account access, announcements, notifications, job information, and updated record views. Administrators provide maintenance inputs such as alumni record management, job posting actions, announcement publishing, and report requests."
    )
    add_paragraph(
        document,
        "The system processes these exchanges and returns outputs such as dashboards, managed records, submission results, communication updates, and analytical summaries. The Level 0 DFD therefore explains the overall system boundary without yet expanding the smaller internal processes."
    )

    document.add_page_break()

    add_chapter_heading(document, "CHAPTER IV", "METHODOLOGY")
    add_paragraph(
        document,
        "This study uses the Waterfall Model as the development methodology. The Waterfall Model was chosen because it provides a clear and organized sequence of phases, which is suitable for a capstone project that requires formal documentation, design outputs, implementation, testing, and evaluation."
    )

    phases = [
        (
            "4.1 Requirements Gathering",
            "During this phase, the development team identified the core needs of the Alumni Relation and Linkages Office. The team examined the current way alumni records, employment data, announcements, feedback, and tracer-related processes were being handled and translated these observations into system requirements.",
        ),
        (
            "4.2 System Design",
            "The system design phase focused on the structure of the database, the layout of the alumni and administrator interfaces, and the flow of information between modules. This phase produced design artifacts such as the ERD, DFD Level 0, module breakdowns, and process-oriented planning for registration, tracer management, and reporting.",
        ),
        (
            "4.3 Implementation",
            "Implementation involved building the actual web application using Laravel, PHP, PostgreSQL, Blade, Tailwind CSS, Vite, and Chart.js. Modules were created for authentication, profile maintenance, tracer records, jobs, feedback, announcements, notifications, and analytics.",
        ),
        (
            "4.4 Testing",
            "After development, the system was tested to verify that each module worked as intended. Feature tests were used to check important workflows such as login, account verification, password reset, record management, feedback submission, and other core operations.",
        ),
        (
            "4.5 Deployment",
            "For deployment, the application configuration, migrations, seed data, and environment settings were prepared so that the system could run in a controlled local or hosted environment. This phase ensures that the developed platform can be used beyond development mode.",
        ),
        (
            "4.6 Maintenance",
            "Maintenance covers future updates, refinements, and corrections after deployment. Although long-term maintenance is beyond the immediate implementation phase, the project includes documentation and testing support that can help future developers or administrators continue improving the system.",
        ),
    ]
    for heading, body in phases:
        add_section_heading(document, heading)
        add_paragraph(document, body)

    add_section_heading(document, "Use Cases for Alumni and Administrator")
    add_paragraph(
        document,
        "For alumni, the main use cases include registration, login, viewing dashboard information, maintaining profile records, submitting tracer data, browsing jobs, applying for opportunities, reading announcements, receiving notifications, and submitting feedback."
    )
    add_paragraph(
        document,
        "For administrators, the main use cases include login, managing alumni records, reviewing job submissions, publishing announcements, monitoring feedback, viewing analytics and reports, and maintaining system activity records. These use cases reflect the major interactions that the system was designed to support."
    )

    add_section_heading(document, "Methodology Summary")
    add_paragraph(
        document,
        "Through the Waterfall Model, the development of the proposed system follows a structured path from requirements gathering up to testing and maintenance. This makes the methodology suitable for a formal academic study where each phase needs to be documented and justified clearly."
    )

    add_paragraph(document, "REFERENCES", alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    for reference_text in REFERENCES:
        add_paragraph(document, reference_text)

    return document


def main() -> None:
    output_path = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_OUTPUT_PATH
    document = build_document()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    print(f"Created: {output_path}")


if __name__ == "__main__":
    main()
