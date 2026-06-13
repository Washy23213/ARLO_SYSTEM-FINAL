import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SCHEMA_PATH = ROOT / "storage" / "app" / "database_schema_current.json"
OUT_PATH = ROOT / "docs" / "arlo-database-table-dictionary.docx"

ACCENT = "0B3B1F"
ACCENT_2 = "1E7A34"
LIGHT_GREEN = "EAF7EF"
LIGHT_AMBER = "FFF7E0"
LIGHT_GRAY = "F4F6F4"
BORDER = "D7DED8"


TABLE_NOTES = {
    "users": {
        "category": "Core alumni record",
        "story": (
            "The users table is the center of the alumni side of the system. It stores both CSV-imported alumni roster "
            "records and alumni accounts that already completed registration. When an alumnus signs up, the system "
            "checks the submitted name, birthdate, college, graduation year, and optional student number against this "
            "table. After registration, the same record becomes the alumni account used for login, profile management, "
            "notifications, tracer submission, and analytics."
        ),
        "where": "Registration verification, alumni login, alumni directory, profile management, tracer analytics, notifications, and CSV batch import.",
        "how": "Created by CSV import or manual/admin entry; completed by alumni registration; updated by alumni profile forms, admin maintenance, and super admin account correction.",
    },
    "admins": {
        "category": "Administrative account",
        "story": (
            "The admins table stores regular administrator accounts used by ARLO staff. These accounts are separate from alumni "
            "records so office personnel do not mix with graduate records. Admins use this table to log in to the operational "
            "dashboard where they manage alumni records, job postings, announcements, feedback, reports, analytics, and audit logs."
        ),
        "where": "Admin login, dashboard access, alumni management, jobs, announcements, feedback, analytics, reports, and audit logging.",
        "how": "Created and maintained by the Super Admin. Admin actions are linked to this table through related records and system logs.",
    },
    "super_admins": {
        "category": "Super Admin account",
        "story": (
            "The super_admins table stores the highest account-control role. It is separated from the admins table because the "
            "Super Admin has a different interface and responsibility: managing admin accounts and maintaining alumni accounts, "
            "not handling daily operational modules such as announcements or analytics."
        ),
        "where": "Super Admin login, administrator account creation, administrator account editing, role/access control, and alumni account correction.",
        "how": "Usually seeded or created by a trusted setup process; used to create, update, or delete admin accounts and recorded in system logs.",
    },
    "roles": {
        "category": "Role definition",
        "story": (
            "The roles table stores the administrator role names used by the role-based access control design. Instead of treating every "
            "regular administrator the same way, the system can classify an admin account as Administrator, Alumni Records Manager, or Job Manager."
        ),
        "where": "Super Admin account management, admin login authorization, scoped administrator navigation, and permission checks.",
        "how": "Seeded by the RBAC migration and synchronized when the Super Admin creates or updates an administrator account.",
    },
    "permissions": {
        "category": "Permission definition",
        "story": (
            "The permissions table stores the individual actions that can be granted to administrator roles. These records describe access such as "
            "managing alumni, managing jobs, managing announcements, handling feedback, viewing analytics, and viewing audit logs."
        ),
        "where": "Admin route middleware, module access checks, role setup, and Super Admin account-control workflows.",
        "how": "Seeded by the RBAC migration and connected to roles through the role_permission table.",
    },
    "admin_role": {
        "category": "Role assignment",
        "story": (
            "The admin_role table connects administrator accounts to their assigned roles. It is a pivot table that allows the system to answer "
            "which administrator has which role without storing all access rules directly in the admins table."
        ),
        "where": "Administrator account creation, administrator role editing, login access checks, and scoped module access.",
        "how": "Created or updated when an administrator account is saved and the selected role is synchronized.",
    },
    "role_permission": {
        "category": "Permission assignment",
        "story": (
            "The role_permission table connects each role to the permissions it grants. Through this pivot table, the system can check whether an "
            "administrator is allowed to open a module or perform an administrative action."
        ),
        "where": "Permission middleware, admin navigation, scoped administrator workflows, and RBAC maintenance.",
        "how": "Seeded by the RBAC migration according to the default permission set for each administrator role.",
    },
    "tracer_records": {
        "category": "Tracer and employment data",
        "story": (
            "The tracer_records table is where the graduate tracer story becomes measurable. After an alumnus completes tracer forms, "
            "the answers are stored here and connected to the alumni account. The analytics dashboard reads this table to summarize "
            "employment status, industry distribution, job alignment, satisfaction, salary range, and related graduate outcomes."
        ),
        "where": "Tracer form, employment analytics, report generation, charts, graduate outcome monitoring, and institutional planning.",
        "how": "Created or updated when alumni submit tracer information. Administrators read it for monitoring and reporting.",
    },
    "jobs": {
        "category": "Published job opportunity",
        "story": (
            "The jobs table stores employment opportunities that are visible to alumni. In the current implementation, alumni can browse "
            "job postings but do not apply through the system. This keeps the module as an employment information board managed by the office."
        ),
        "where": "Job browsing page, admin job management, dashboard counts, and employment opportunity publication.",
        "how": "Created and edited by administrators or produced from approved alumni job submissions; deleted directly from the database when removed.",
    },
    "job_submissions": {
        "category": "Alumni-submitted job opportunity",
        "story": (
            "The job_submissions table stores job opportunities suggested by alumni. It acts as a review lane: alumni can notify ARLO about "
            "an opportunity, and administrators decide whether to approve or reject it. Approved submissions can become published job records."
        ),
        "where": "Alumni job opportunity submission, admin job submission review, approval/rejection notifications, and audit trail.",
        "how": "Created by alumni; reviewed by admins; linked to users as submitters and admins as reviewers.",
    },
    "announcements": {
        "category": "Communication record",
        "story": (
            "The announcements table stores official ARLO updates. It connects the previous Facebook-centered communication process to the "
            "web system by allowing administrators to create and publish announcements in one place, with optional Facebook synchronization."
        ),
        "where": "Admin announcement management, alumni announcements, notification generation, Facebook sync, and announcement views.",
        "how": "Created and edited by administrators; published to alumni; deleted permanently when the admin chooses Delete.",
    },
    "announcement_reads": {
        "category": "Read tracking",
        "story": (
            "The announcement_reads table answers a simple but useful question: has this alumnus already opened this announcement? "
            "It supports read/unread indicators so the system can make announcements easier to monitor."
        ),
        "where": "Alumni announcement detail, notification indicators, read/unread status, and announcement engagement monitoring.",
        "how": "Created when an alumnus opens or reads an announcement; deleted automatically when the related announcement is deleted.",
    },
    "notifications": {
        "category": "System notification",
        "story": (
            "The notifications table is the system's message inbox. It stores notices, reminders, warnings, announcement alerts, "
            "and job submission review results for registered alumni. This supports direct system-based communication instead of relying only on Facebook or manual messages."
        ),
        "where": "Alumni notification page, admin warning/reminder/notice buttons, announcement alerts, and job submission results.",
        "how": "Created by Laravel notification classes or admin actions; marked read when alumni open or acknowledge them.",
    },
    "feedback": {
        "category": "Alumni feedback",
        "story": (
            "The feedback table stores alumni concerns, suggestions, and inquiries. It gives ARLO a structured place to receive and review alumni feedback "
            "instead of keeping concerns scattered across informal messages."
        ),
        "where": "Alumni feedback form, admin feedback management, response tracking, and dashboard summaries.",
        "how": "Created by alumni; reviewed and updated by administrators.",
    },
    "system_logs": {
        "category": "Audit trail",
        "story": (
            "The system_logs table is the accountability record of the platform. Important actions such as logins, record changes, admin account changes, "
            "announcements, notifications, and deletes are recorded here so the office can trace what happened and who performed the action."
        ),
        "where": "Audit log module, security monitoring, accountability checks, and administrative review.",
        "how": "Created automatically by audit logging services whenever important system actions occur.",
    },
    "sessions": {
        "category": "Authentication support",
        "story": (
            "The sessions table supports active logins when the database session driver is used. It stores technical session payloads so the system can remember "
            "that a user is already logged in while moving across protected pages."
        ),
        "where": "Login state, protected route access, session management, and logout behavior.",
        "how": "Created when a user logs in or starts a session; removed or expired when the session ends.",
    },
    "password_reset_tokens": {
        "category": "Account recovery support",
        "story": (
            "The password_reset_tokens table stores temporary reset tokens. It only receives data when a user requests a password reset link. "
            "The record is short-lived and is used to confirm that the reset request is valid."
        ),
        "where": "Forgot password and reset password workflow for alumni, admins, and Super Admin accounts.",
        "how": "Created when a reset link is requested; replaced by newer reset requests and consumed or expired after password reset.",
    },
    "cache": {
        "category": "Laravel support",
        "story": (
            "The cache table stores temporary system values that Laravel can reuse. It is not a direct alumni feature, but it helps the application avoid repeated work "
            "for values that can be safely stored for a short time."
        ),
        "where": "Framework-level caching, temporary lookups, and performance support.",
        "how": "Created and cleared automatically by Laravel cache operations.",
    },
    "cache_locks": {
        "category": "Laravel support",
        "story": (
            "The cache_locks table prevents two processes from performing the same sensitive cached operation at the same time. It is a support table for safe background or concurrent operations."
        ),
        "where": "Framework-level lock management and concurrency protection.",
        "how": "Created temporarily when Laravel obtains a database cache lock; removed when the lock expires or is released.",
    },
    "failed_jobs": {
        "category": "Laravel support",
        "story": (
            "The failed_jobs table records background jobs that fail. In the current system the queue driver is set to sync and queue_jobs/job_batches were removed, "
            "so this table is normally empty unless queued processing is reintroduced later."
        ),
        "where": "Framework-level troubleshooting for failed queued tasks.",
        "how": "Receives records only when a queued job fails while a queue worker is being used.",
    },
    "migrations": {
        "category": "Laravel support",
        "story": (
            "The migrations table is Laravel's memory of database changes. Every time a migration runs, Laravel records it here so the system knows which table changes "
            "have already been applied."
        ),
        "where": "Database version control, deployment, migration status checks, and schema maintenance.",
        "how": "Updated automatically by Laravel when php artisan migrate is executed.",
    },
}


REMOVED_TABLES = [
    ("job_applications", "Removed from the current implementation because alumni do not apply for jobs through the portal. Jobs are used for browsing and information only."),
    ("queue_jobs", "Removed because the project now uses sync queue execution instead of storing queued jobs in the database."),
    ("job_batches", "Removed because Laravel batch queue processing is not used in the current system."),
    ("job_matches", "Not present in the current database schema."),
]

REMOVED_FIELDS = [
    ("jobs.salary_min", "Removed because the maintained job salary value is the selected salary_range label."),
    ("jobs.salary_max", "Removed because it was only a derived copy of the salary_range label and was not needed for the current workflow."),
    ("jobs.applications_count", "Removed with the direct job-application workflow. The current system uses job_submissions for alumni-suggested opportunities."),
]


def normalize_type(col):
    data_type = col.get("data_type") or ""
    max_len = col.get("character_maximum_length")
    precision = col.get("numeric_precision")
    scale = col.get("numeric_scale")

    if data_type == "character varying" and max_len:
        return f"varchar({max_len})"
    if data_type == "numeric" and precision:
        return f"numeric({precision},{scale or 0})"
    if data_type == "timestamp without time zone":
        return "timestamp"
    if data_type == "timestamp with time zone":
        return "timestampTz"
    if data_type == "USER-DEFINED":
        return "custom"
    return data_type


def title_from_name(name):
    return name.replace("_", " ").title()


def column_description(table, col):
    name = col["column_name"]
    pretty = title_from_name(name).lower()

    custom = {
        "id": "Primary identifier for the record.",
        "created_at": "Date and time when the record was created.",
        "updated_at": "Date and time when the record was last updated.",
        "deleted_at": "Date and time used for soft deletion when applicable.",
        "email": "Email address used for contact, login, or account recovery.",
        "password": "Encrypted password used for secure authentication.",
        "remember_token": "Token used by Laravel for remember-me authentication.",
        "email_verified_at": "Date and time when the email address was verified.",
        "firstname": "First name of the account holder.",
        "lastname": "Last name of the account holder.",
        "middlename": "Middle name or middle initial of the account holder.",
        "phone": "Contact phone number.",
        "role": "Access role used by the application to control permissions.",
        "user_id": "Identifier of the alumni user connected to this record.",
        "admin_id": "Identifier of the administrator connected to this action.",
        "super_admin_id": "Identifier of the Super Admin connected to this action.",
        "title": "Display title of the record.",
        "content": "Main body or message content.",
        "status": "Current workflow state of the record.",
        "type": "Classification or category of the record.",
        "priority": "Importance level used to organize or highlight the record.",
        "name": "System name used internally by the application.",
        "label": "User-readable label displayed in the system.",
        "guard_name": "Authentication guard where the role or permission applies.",
        "read_at": "Date and time when the notification or announcement was read.",
        "payload": "Serialized technical session or job data used by Laravel.",
        "token": "Secure token used for password reset validation.",
        "migration": "Name of a Laravel migration file that has already run.",
        "batch": "Migration batch number used by Laravel.",
        "action": "Name of the system action that was recorded.",
        "description": "Explanation, details, or description for the record.",
        "ip_address": "IP address captured during the action or session.",
        "user_agent": "Browser or device information captured during the action or session.",
        "metadata": "Additional structured information stored with the record.",
        "data": "Serialized notification payload shown to the user.",
        "notifiable_type": "Laravel notification target model type.",
        "notifiable_id": "Identifier of the notification recipient.",
        "queue": "Queue name used by Laravel background job support.",
        "exception": "Error details captured when a background job fails.",
        "failed_at": "Date and time when the background job failed.",
        "key": "Cache or lock key used by Laravel.",
        "value": "Cached value stored by Laravel.",
        "expiration": "Unix timestamp indicating when the cache value expires.",
        "owner": "Owner token of a Laravel cache lock.",
        "created_by": "Administrator who created the record.",
        "archived_by": "Administrator who archived the record before archive behavior was replaced by direct delete.",
        "archived_at": "Previous archive timestamp column retained in the schema, although the current delete action removes records directly.",
    }

    if name in custom:
        return custom[name]

    table_specific = {
        ("users", "alumni_idnum"): "Official alumni ID number used for graduate reference and verification.",
        ("users", "student_idnum"): "Student ID number used as an optional additional verification field.",
        ("users", "birthdate"): "Birth date used to strengthen alumni registration verification.",
        ("users", "department"): "College or department graduated from.",
        ("users", "degree"): "Degree or program completed by the alumnus.",
        ("users", "dt_grad"): "Exact graduation date when available from the alumni record.",
        ("users", "year_graduated"): "Graduation year used for filtering, reports, and verification.",
        ("users", "is_batch_imported"): "Marks records imported from the CSV graduate reference file before full account registration.",
        ("users", "batch_imported_at"): "Date and time when the alumni roster record was imported.",
        ("users", "profile_visibility"): "Controls whether selected alumni profile information can be displayed.",
        ("users", "data_privacy_consent"): "Stores whether the alumnus agreed to the system's data privacy consent.",
        ("users", "newsletter_opt_in"): "Stores whether the alumnus agreed to receive announcement and newsletter emails.",
        ("tracer_records", "tracer_number"): "Human-readable tracer reference number used in reports.",
        ("tracer_records", "tracer_date"): "Date when the tracer record was submitted or updated.",
        ("tracer_records", "employment_status"): "Employment status used in tracer analytics, such as employed, self-employed, or unemployed.",
        ("tracer_records", "employer_name"): "Name of the employer reported by the alumnus.",
        ("tracer_records", "job_title"): "Job position or occupation reported by the alumnus.",
        ("tracer_records", "employment_proof_path"): "Stored file path of the uploaded employment proof document.",
        ("tracer_records", "employment_proof_status"): "Review status of the uploaded employment proof.",
        ("tracer_records", "employment_proof_reviewed_by_guard"): "Authentication guard of the reviewer who evaluated the employment proof.",
        ("tracer_records", "employment_proof_reviewed_by_id"): "Identifier of the reviewer who evaluated the employment proof.",
        ("tracer_records", "employment_proof_rejection_note"): "Reason recorded when an employment proof upload is rejected.",
        ("admin_role", "admin_id"): "Administrator account assigned to the role.",
        ("admin_role", "role_id"): "Role assigned to the administrator account.",
        ("role_permission", "role_id"): "Role that receives the permission.",
        ("role_permission", "permission_id"): "Permission granted to the role.",
        ("job_submissions", "reviewed_by"): "Administrator who reviewed the alumni-submitted job opportunity.",
        ("job_submissions", "reviewed_at"): "Date and time when the submitted job opportunity was reviewed.",
        ("job_submissions", "admin_notes"): "Notes entered by the administrator during review.",
        ("jobs", "salary_range"): "Salary range label displayed with the published job posting.",
        ("jobs", "is_active"): "Indicates whether the job posting is visible and active.",
        ("jobs", "views_count"): "Number of times the job posting was viewed.",
        ("announcements", "is_published"): "Indicates whether the announcement is visible to alumni.",
        ("announcements", "published_at"): "Date and time when the announcement was published.",
        ("announcements", "views_count"): "Number of times the announcement was viewed.",
        ("announcements", "facebook_post_id"): "Facebook post identifier used when syncing with the official page.",
        ("announcements", "facebook_post_url"): "URL of the related Facebook post.",
        ("announcement_reads", "announcement_id"): "Announcement record that was read by the alumnus.",
        ("feedback", "message"): "Feedback, concern, or inquiry submitted by an alumnus.",
        ("feedback", "response"): "Administrator response to the submitted feedback.",
    }

    if (table, name) in table_specific:
        return table_specific[(table, name)]

    if name.startswith("is_"):
        return f"Boolean flag indicating whether {pretty.replace('is ', '')} applies."
    if name.endswith("_at"):
        return f"Date and time related to {pretty.replace(' at', '')}."
    if name.endswith("_date"):
        return f"Date value for {pretty.replace(' date', '')}."
    if name.endswith("_url"):
        return f"URL reference for {pretty.replace(' url', '')}."
    if name.endswith("_path"):
        return f"Stored file path for {pretty.replace(' path', '')}."
    if name.endswith("_count"):
        return f"Numeric counter for {pretty.replace(' count', '')}."
    if name.endswith("_id"):
        return f"Identifier connected to {pretty.replace(' id', '')}."

    return f"Stores the {pretty} value used by the {table} record."


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            cell = row.cells[idx]
            cell.width = Pt(width / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def style_table(table, widths, header_fill=LIGHT_GREEN, font_size=8.5):
    set_table_width(table, widths)
    for row_idx, row in enumerate(table.rows):
        set_cant_split(row)
        if row_idx == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    if row_idx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(11, 59, 31)


def add_table_row(table, values):
    row = table.add_row().cells
    for cell, value in zip(row, values):
        cell.text = str(value or "")
    return row


def add_callout(doc, title, body, fill=LIGHT_AMBER):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="E2C76D")
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(11, 59, 31)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for run in p2.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)
    style_table(table, [9360], header_fill=fill, font_size=10)
    doc.add_paragraph()


def add_small_metadata_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Item"
    hdr[1].text = "Value"
    for key, value in rows:
        add_table_row(table, [key, value])
    style_table(table, [2400, 6960], header_fill=LIGHT_GREEN, font_size=9)
    doc.add_paragraph()


def add_summary_table(doc, rows):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Table", "Category", "Rows", "Main Purpose", "Main Users"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for row in rows:
        add_table_row(table, row)
    style_table(table, [1700, 1700, 800, 3360, 1800], header_fill=LIGHT_GREEN, font_size=8.3)
    doc.add_paragraph()


def add_removed_items_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for row in rows:
        add_table_row(table, row)
    style_table(table, [2300, 7060], header_fill=LIGHT_AMBER, font_size=9)
    doc.add_paragraph()


def build_constraint_maps(constraints):
    primary = defaultdict(set)
    unique = defaultdict(set)
    foreign = defaultdict(list)

    for c in constraints:
        table = c["table_name"]
        col = c.get("column_name")
        ctype = c["constraint_type"]
        if not col:
            continue
        if ctype == "PRIMARY KEY":
            primary[table].add(col)
        elif ctype == "UNIQUE":
            unique[table].add(col)
        elif ctype == "FOREIGN KEY":
            foreign[table].append((col, c.get("foreign_table_name"), c.get("foreign_column_name")))

    return primary, unique, foreign


def relationship_text(table, col, primary, unique, foreign):
    parts = []
    if col in primary.get(table, set()):
        parts.append("PK")
    if col in unique.get(table, set()):
        parts.append("Unique")
    for fk_col, fk_table, fk_column in foreign.get(table, []):
        if fk_col == col:
            parts.append(f"FK -> {fk_table}.{fk_column}")
    return "; ".join(parts) if parts else "-"


def clean_default(value):
    if not value:
        return ""
    value = str(value)
    value = re.sub(r"::[\w\s]+", "", value)
    return value


def add_field_dictionary(doc, table_name, columns, primary, unique, foreign):
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Field", "Data Type", "Required", "Key / Relationship", "Purpose in the System"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header

    for col in columns:
        required = "Yes" if col.get("is_nullable") == "NO" else "No"
        default = clean_default(col.get("column_default"))
        if default:
            required = f"{required}; default: {default}"
        add_table_row(table, [
            col["column_name"],
            normalize_type(col),
            required,
            relationship_text(table_name, col["column_name"], primary, unique, foreign),
            column_description(table_name, col),
        ])

    style_table(table, [1550, 1450, 1500, 2050, 2810], header_fill=LIGHT_GREEN, font_size=8.0)
    doc.add_paragraph()


def add_table_section(doc, table_name, row_count, columns, primary, unique, foreign):
    note = TABLE_NOTES[table_name]
    h = doc.add_heading(f"{table_name}", level=2)
    h.runs[0].font.color.rgb = RGBColor(11, 59, 31)

    p = doc.add_paragraph()
    p.add_run("Purpose. ").bold = True
    p.add_run(note["story"])

    p = doc.add_paragraph()
    p.add_run("Where it appears in the system. ").bold = True
    p.add_run(note["where"])

    p = doc.add_paragraph()
    p.add_run("How records are created or changed. ").bold = True
    p.add_run(note["how"])

    rels = foreign.get(table_name, [])
    if rels:
        rel_text = "; ".join(f"{col} -> {ft}.{fc}" for col, ft, fc in rels)
    else:
        rel_text = "No direct foreign key from this table in the current schema."

    add_small_metadata_table(doc, [
        ("Table category", note["category"]),
        ("Current development row count", row_count),
        ("Direct relationships", rel_text),
    ])
    add_field_dictionary(doc, table_name, columns, primary, unique, foreign)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Title", 22, ACCENT),
        ("Subtitle", 11, "42574A"),
        ("Heading 1", 16, ACCENT),
        ("Heading 2", 13, ACCENT),
        ("Heading 3", 11, ACCENT_2),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name.startswith("Heading") or style_name == "Title"


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = "CMU OARL - Database Table Dictionary"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(85, 102, 92)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    p._p.append(field)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(85, 102, 92)


def main():
    data = json.loads(SCHEMA_PATH.read_text(encoding="utf-8-sig"))
    tables = [t["table_name"] for t in data["tables"] if t["table_name"] in TABLE_NOTES]
    columns_by_table = defaultdict(list)
    for col in data["columns"]:
        if col["table_name"] in TABLE_NOTES:
            columns_by_table[col["table_name"]].append(col)

    primary, unique, foreign = build_constraint_maps(data["constraints"])

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    add_header_footer(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Database Table Dictionary")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("A Web-Based Tracer and Employment Analytics System for the Alumni Relations and Linkages Office")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Central Mindanao University | May 2026").bold = True

    add_small_metadata_table(doc, [
        ("Prepared for", "WASHELEY TAN and ELVIN LAR GONZALES"),
        ("Program", "Bachelor of Science in Information Technology"),
        ("Database verified from", f"{data['connection']} connection, database {data['database']}"),
        ("Schema generated at", data["generated_at"]),
        ("Current table count", len(tables)),
    ])

    add_callout(
        doc,
        "Verification Note",
        "This document was generated from the current Laravel and PostgreSQL database schema. "
        f"The live database contains {len(tables)} documented tables. It does not contain job_applications, queue_jobs, job_batches, or job_matches in the current implementation. "
        "Analytics is computed from existing operational tables, especially users and tracer_records; it does not use a separate analytics table.",
        LIGHT_AMBER,
    )

    doc.add_heading("1. How the Database Works as a Story", level=1)
    paragraphs = [
        (
            "The database begins with accounts. Alumni records are stored in the users table, regular administrators are stored in the admins table, "
            "and the higher account-control role is stored in the super_admins table. This separation keeps graduate records away from staff accounts and "
            "makes the Super Admin role clearer in the system design."
        ),
        (
            "When ARLO imports a CSV file, the imported alumni are stored in users as roster records. They are not treated as fully registered accounts yet. "
            "During registration, the alumnus must match the stored name, birthdate, college, graduation year, and optional student number. Once verified, "
            "the same record becomes the account used for login and profile completion."
        ),
        (
            "After login, the alumnus can maintain profile information and submit tracer data. The profile remains in users, while detailed tracer answers are stored "
            "in tracer_records. This is the main data source for charts, reports, and employment analytics."
        ),
        (
            "For communication, administrators create announcements in announcements, while read tracking is stored in announcement_reads. Direct reminders, warnings, "
            "and system messages are stored in notifications. Feedback from alumni is stored in feedback so administrators can review and respond in an organized way."
        ),
        (
            "For employment support, administrators manage published job opportunities in jobs. Alumni can notify ARLO about job opportunities through job_submissions. "
            "The current system does not store job applications because alumni do not apply for jobs through the portal."
        ),
        (
            "Finally, system_logs records important actions for accountability. Laravel support tables such as migrations, sessions, cache, cache_locks, password_reset_tokens, "
            "and failed_jobs support framework operations such as schema tracking, login sessions, recovery links, temporary storage, and troubleshooting."
        ),
    ]
    for text in paragraphs:
        doc.add_paragraph(text)

    doc.add_heading("2. Current Table Summary", level=1)
    summary_rows = []
    for table in tables:
        note = TABLE_NOTES[table]
        users = {
            "Core alumni record": "Alumni, Admin, Super Admin",
            "Administrative account": "Administrator, Super Admin",
            "Super Admin account": "Super Admin",
            "Role definition": "Super Admin, Administrator",
            "Permission definition": "Super Admin, Administrator",
            "Role assignment": "Super Admin, System",
            "Permission assignment": "System",
            "Tracer and employment data": "Alumni, Administrator",
            "Published job opportunity": "Alumni, Administrator",
            "Alumni-submitted job opportunity": "Alumni, Administrator",
            "Communication record": "Alumni, Administrator",
            "Read tracking": "Alumni, System",
            "System notification": "Alumni, Administrator, System",
            "Alumni feedback": "Alumni, Administrator",
            "Audit trail": "Administrator, Super Admin",
            "Authentication support": "System",
            "Account recovery support": "System",
            "Laravel support": "System",
        }.get(note["category"], "System")
        summary_rows.append([
            table,
            note["category"],
            data["row_counts"].get(table, 0),
            note["where"].split(",")[0],
            users,
        ])
    add_summary_table(doc, summary_rows)

    doc.add_heading("3. Tables Not Present in the Current Database", level=1)
    doc.add_paragraph(
        "The following names appeared during design discussion or earlier development stages, but they are not part of the current live schema."
    )
    add_removed_items_table(doc, ["Table Name", "Reason"], REMOVED_TABLES)

    doc.add_heading("4. Fields Removed from the Current Schema", level=1)
    doc.add_paragraph(
        "The following fields were removed from the ERD because they were obsolete or derived from fields already maintained by the system."
    )
    add_removed_items_table(doc, ["Removed Field", "Reason"], REMOVED_FIELDS)

    doc.add_heading("5. Detailed Table Dictionary", level=1)
    for table in tables:
        add_table_section(
            doc,
            table,
            data["row_counts"].get(table, 0),
            columns_by_table[table],
            primary,
            unique,
            foreign,
        )

    doc.add_heading("6. Notes for the Manuscript", level=1)
    notes = [
        "If the paper still mentions direct job applications, update that part to match the current implementation: alumni browse job postings and submit job opportunities for review, but they do not apply for jobs inside the portal.",
        "If the ERD is updated, it should show users, admins, and super_admins as separate account tables.",
        "The analytics module should be described as a computed feature that reads from users, tracer_records, jobs, job_submissions, announcements, feedback, notifications, and system_logs, not as a separate analytics table.",
        "The queue_jobs and job_batches tables should not be included in the final ERD because they were removed from the current database.",
    ]
    for item in notes:
        doc.add_paragraph(item, style=None)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
