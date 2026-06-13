from __future__ import annotations

import json
from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "arlo-results-discussion-ui-modules.docx"
SCHEMA_JSON = ROOT / "storage" / "app" / "database_schema_current.json"

SKILL_DIR = Path(
    r"C:\Users\washe\.codex\plugins\cache\openai-primary-runtime\documents\26.430.10722\skills\documents"
)
sys.path.insert(0, str(SKILL_DIR / "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights  # type: ignore


ACCENT = RGBColor(0x00, 0x67, 0x38)
DARK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x56, 0x69)
LIGHT_GREEN = "EAF8F0"
LIGHT_BLUE = "EEF5FF"
LIGHT_AMBER = "FFF7E6"
LIGHT_RED = "FFF1F2"
GRID = "D9E2DE"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None, size: int = 10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def style_table(table, header_fill: str = "F7FAF8") -> None:
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = DARK
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.name = "Arial"
                    if run.font.size is None:
                        run.font.size = Pt(10)


def add_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run("CMU OARL Results and Discussion Guide")
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        p.paragraph_format.space_after = Pt(2)

        p_pr = p._p.get_or_add_pPr()
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "D9E2DE")
        borders.append(bottom)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.text = "Page "
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        fp._p.append(fld_begin)
        fp._p.append(instr)
        fp._p.append(fld_end)
        for run in fp.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.color.rgb = MUTED


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after, color in [
        ("Title", 22, 0, 12, ACCENT),
        ("Subtitle", 12, 0, 12, MUTED),
        ("Heading 1", 16, 12, 6, ACCENT),
        ("Heading 2", 14, 10, 4, DARK),
        ("Heading 3", 12, 8, 3, DARK),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = name.startswith("Heading") or name == "Title"
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_para(doc: Document, text: str, style: str | None = None, bold_label: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    if bold_label and text.startswith(bold_label):
        r = p.add_run(bold_label)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11)
        p.add_run(text[len(bold_label):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Arial"
        if run.font.size is None:
            run.font.size = Pt(11)


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_GREEN) -> None:
    table = doc.add_table(rows=1, cols=1)
    apply_table_geometry(table, [9360], table_width_dxa=9360, indent_dxa=0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT if fill != LIGHT_RED else RGBColor(0xBE, 0x12, 0x3C)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for run in p2.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = DARK
    doc.add_paragraph()


def add_metadata_table(doc: Document) -> None:
    rows = [
        ("Project", "A Web-Based Tracer and Employment Analytics System for the Alumni Relations and Linkages Office"),
        ("Institution", "Central Mindanao University"),
        ("Prepared For", "Chapter V: Results and Discussion"),
        ("Prepared Date", "May 2026"),
        ("Implementation Checked", "Laravel 12 / PHP 8.2 / PostgreSQL current workspace"),
    ]
    table = doc.add_table(rows=1, cols=2)
    widths = column_widths_from_weights(9360, [1.4, 4.6])
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=0)
    set_cell_text(table.cell(0, 0), "Field", bold=True)
    set_cell_text(table.cell(0, 1), "Details", bold=True)
    for field, detail in rows:
        row = table.add_row().cells
        set_cell_text(row[0], field, bold=True, color=MUTED)
        set_cell_text(row[1], detail)
    style_table(table)


def add_system_verification_table(doc: Document) -> None:
    tables = []
    if SCHEMA_JSON.exists():
        schema = json.loads(SCHEMA_JSON.read_text(encoding="utf-8"))
        tables = [item["table_name"] for item in schema.get("tables", [])]
    else:
        tables = [
            "admins", "announcement_reads", "announcements", "cache", "cache_locks",
            "failed_jobs", "feedback", "job_submissions", "jobs", "migrations",
            "notifications", "password_reset_tokens", "sessions", "super_admins",
            "system_logs", "tracer_records", "users",
        ]

    rows = [
        ("Account structure", "Alumni records are in users; regular administrators are in admins; Super Admin accounts are in super_admins."),
        ("Job feature scope", "Alumni browse job postings and submit job opportunities for admin review. They do not submit job applications through the portal."),
        ("Current database tables", ", ".join(tables)),
        ("Removed or not active", "job_applications, queue_jobs, job_batches, and job_matches are not part of the current live database."),
        ("Analytics storage", "Analytics are computed from operational tables such as users, tracer_records, jobs, job_submissions, feedback, and announcements. There is no separate analytics table."),
        ("Notification behavior", "Admin notices are saved inside the alumni account and are also sent by email when the Gmail/mail configuration is working."),
    ]
    table = doc.add_table(rows=1, cols=2)
    widths = column_widths_from_weights(9360, [1.6, 4.4])
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=0)
    set_cell_text(table.cell(0, 0), "Checked Area", bold=True)
    set_cell_text(table.cell(0, 1), "Current System Result", bold=True)
    for field, detail in rows:
        row = table.add_row().cells
        set_cell_text(row[0], field, bold=True, color=ACCENT)
        set_cell_text(row[1], detail)
    style_table(table, "F4FBF7")


def module_blocks():
    return [
        {
            "title": "5.1 Login and Account Recovery Interface",
            "result": "The system produced a secure entry point where alumni, administrators, and Super Admin users sign in using registered email credentials, while forgotten passwords are handled through email-based reset links.",
            "discussion": "This interface is the first control point of the platform. When a user opens the login page, the system asks for the registered email address and password before allowing access to any protected module. If the user forgets the password, the account recovery screen requests the registered email and creates a password reset token in the database. This supports the study objective on secure registration, authentication, and account recovery because the system does not expose alumni records, tracer forms, analytics, or administrative tools to unauthenticated users.",
            "screenshot": "Capture the login page, then capture the forgot-password page after entering a valid registered email and receiving a professional success message.",
            "connect": "Connects to Objective 1, Functional Requirement: Alumni Authentication and Account Recovery, Administrator Authentication.",
        },
        {
            "title": "5.2 Alumni Registration and Verification Interface",
            "result": "The registration module verifies alumni against the imported alumni roster before account creation and requires email code verification before password setup and final registration.",
            "discussion": "The registration process follows the problem identified in the old workflow: ARLO previously needed to manually check whether a respondent was a legitimate CMU alumnus. In the system, the alumnus first enters identity and academic details that must match the roster imported from CSV. The birthdate field strengthens verification because it adds a personal data point that is not visible from the name alone. After the alumni record is matched, the user must verify the email address through a code sent to Gmail or the configured mail service. This gives the system a stronger assurance that the alumni account belongs to a reachable email address before allowing the user to proceed.",
            "screenshot": "Capture the registration verification step, the email code panel, and the password setup area after successful code verification.",
            "connect": "Connects to Objective 1 and Business Rule: Alumni Verification and Authentication Requirement.",
        },
        {
            "title": "5.3 Alumni Dashboard Interface",
            "result": "The alumni dashboard gives registered alumni a single overview of profile completion, tracer status, recent employment record, notifications, and available portal actions.",
            "discussion": "After login, the alumni dashboard works like the home base of the alumni account. Instead of making the alumnus search through many pages, it summarizes what still needs attention and what has already been submitted. The checklist guides the alumnus toward completing profile and tracer information, while the latest employment record shows the most recent tracer output. This supports usability because the user can immediately understand whether the account is complete, whether tracer information is available, and where to continue next.",
            "screenshot": "Capture the full alumni dashboard with the profile status cards, tracer status, latest employment record, checklist, and recent notifications.",
            "connect": "Connects to Objective 2 and Non-Functional Requirement: Usability.",
        },
        {
            "title": "5.4 Profile and Settings Interface",
            "result": "The profile and settings modules allow alumni to maintain core account details, contact information, academic background, password settings, and account controls.",
            "discussion": "The profile page separates basic alumni identity details from tracer employment records. This prevents the profile page from becoming overloaded and keeps employment updates inside the tracer module. The settings page supports password changes, device/session actions, and account maintenance. Together, these pages respond to ARLO's need for records that can be updated directly by alumni instead of being repeatedly edited in spreadsheets by office personnel.",
            "screenshot": "Capture the profile update form, profile completeness checklist, and account settings/password section.",
            "connect": "Connects to Objective 2, Functional Requirement: Alumni Profile Management, and Business Rule: Profile Information Validation.",
        },
        {
            "title": "5.5 Tracer Form and Employment Record Interface",
            "result": "The tracer module collects graduate outcome information, including employment status, current work details, salary range, industry, skills, further study, satisfaction, and supporting employment proof when required.",
            "discussion": "This module is the center of the tracer system. It changes the previous Google Form workflow into a connected database process where tracer responses belong to a specific registered alumni account. The fields are structured so that the information can later be summarized in analytics, filtered in reports, and reviewed by administrators. When alumni submit tracer information, the system can update both the user profile summary and the tracer_records table, making the data useful for institutional reports and employment analytics.",
            "screenshot": "Capture the tracer form sections for employment status, company/position, salary or income, industry, proof upload, satisfaction, and final submission.",
            "connect": "Connects to Objective 2 and Objective 5, Functional Requirement: Tracer Form Submission.",
        },
        {
            "title": "5.6 Job Opportunities and Alumni Job Notice Interface",
            "result": "The current job module allows alumni to view job opportunities and notify OARL about job openings, but it no longer allows alumni to apply for jobs inside the portal.",
            "discussion": "The job module now acts as an employment information channel rather than a recruitment application system. Alumni can browse approved job postings and submit job opportunities for ARLO review. The submitted opportunity is stored in job_submissions and remains pending until an administrator approves or rejects it. If approved, the opportunity can become a published job posting in the jobs table. This design is cleaner for the study because ARLO manages information about opportunities without becoming the official employer application handler.",
            "screenshot": "Capture the alumni Jobs page, the Notify Job Opportunity button, the job submission modal, and the My Recent Submissions table.",
            "connect": "Connects to Objective 3 after revising the wording from job applications to job opportunity submission.",
        },
        {
            "title": "5.7 Announcements and Notification Interface",
            "result": "The system displays official announcements and in-system notifications, while published announcements and direct admin notices can also be sent by email when mail settings are working.",
            "discussion": "The announcement and notification interface improves communication between ARLO and alumni. Alumni can view portal announcements, Facebook-synced announcement records, and notification details from one place. Read and unread states help users identify which announcements or notices still need attention. This directly responds to the previous reliance on Facebook and manual communication by adding a system-based communication channel that can still connect with Facebook when configured.",
            "screenshot": "Capture the alumni notification list, an announcement detail page, and a read/unread announcement badge.",
            "connect": "Connects to Objective 3, Business Rule: Notification Handling Rule, and Needs Assessment Question 5.",
        },
        {
            "title": "5.8 Feedback Submission Interface",
            "result": "The feedback module gives alumni a structured way to send concerns, suggestions, inquiries, and employability-related comments to ARLO.",
            "discussion": "Instead of receiving feedback only through informal messages, the system provides a feedback form with categories and message fields. This makes each concern easier to store, review, and respond to. For alumni, the page feels like a direct channel to the office. For administrators, the submitted feedback becomes a manageable record inside the database rather than a scattered message outside the system.",
            "screenshot": "Capture the feedback category selection, the message form, and the submit confirmation after sending feedback.",
            "connect": "Connects to Objective 3 and Functional Requirement: Feedback Submission and Feedback Management.",
        },
        {
            "title": "5.9 Admin Dashboard Overview",
            "result": "The admin dashboard summarizes current operational counts such as alumni records, registered alumni, not registered alumni, active jobs, announcements, and monthly activity.",
            "discussion": "The overview is important because it prevents confusion between CSV-imported roster records and actual registered alumni accounts. For example, if 450 alumni are imported from CSV but only 100 have registered, the dashboard can show registered and not registered counts separately. This helps administrators understand why analytics graphs may be empty for a filtered batch even when many roster records exist: CSV-only records are available for verification, but they do not have tracer responses yet.",
            "screenshot": "Capture the admin dashboard stat cards showing Alumni Records, Registered, Not Registered, Active Jobs, Announcements, and This Month.",
            "connect": "Connects to Objective 4 and the recent system decision to separate roster counts from registered alumni counts.",
        },
        {
            "title": "5.10 Alumni Directory, CSV Import, and Record Maintenance",
            "result": "The admin alumni directory now focuses on registered alumni accounts, while CSV import supports roster preparation, signup verification, and comparison against the required template.",
            "discussion": "The alumni directory gives administrators a clean list of alumni accounts that have actually completed registration. This avoids mixing unregistered CSV-only records with active portal users. The CSV import interface helps ARLO upload graduate rosters by comparing the selected file with the required template before import. Extra columns are ignored automatically, but required columns must be present so the system can verify names, IDs, course/program, graduation date, and birthdate consistently.",
            "screenshot": "Capture the Alumni Directory table, the CSV batch import modal, the template preview, the selected file comparison, and the ready/missing column state.",
            "connect": "Connects to Objective 4, Business Rule: Data Consistency Rule, and the alumni verification workflow.",
        },
        {
            "title": "5.11 Direct Alumni Notice, Reminder, and Warning Feature",
            "result": "Administrators and Super Admin users can send a specific notice, reminder, or warning to a registered alumni account using the bell action in the alumni directory.",
            "discussion": "The bell action was changed from a vague button into a specific communication tool. When clicked, it opens a modal where the administrator selects the message type, enters a short title, and writes the full instruction. The notification is saved to the alumni account and is also sent by email when the mail configuration is working. The system blocks CSV-only alumni records from receiving direct notices because those records do not yet represent registered accounts with active portal access.",
            "screenshot": "Capture the alumni directory bell button and the notification modal showing Notice, Reminder, and Warning options.",
            "connect": "Connects to Objective 3 and Business Rule: Notification Handling Rule.",
        },
        {
            "title": "5.12 Job Management and Job Submission Review Interface",
            "result": "The admin job module allows administrators to create, edit, delete, and publish job postings, and to approve or reject alumni-submitted job opportunities.",
            "discussion": "This module gives ARLO control over what employment opportunities appear in the portal. Alumni submissions are not published immediately. They are first reviewed by an administrator, which supports quality control and prevents inaccurate or inappropriate postings from becoming visible. When a submission is approved, the system can create a job record and update the submission status; when rejected, the admin can add notes. Delete actions remove the selected job posting from the database and are recorded through system logs.",
            "screenshot": "Capture the admin Job Management table, the pending submissions section, the approve/reject action, and the delete button.",
            "connect": "Connects to Objective 4 and Business Rule: Job Posting Approval Rule.",
        },
        {
            "title": "5.13 Announcement Management Interface",
            "result": "The announcement management module lets administrators create, edit, publish, delete, filter, and synchronize announcements with Facebook when the integration is configured.",
            "discussion": "The announcement table is designed for office communication. Administrators can filter announcements by keyword, type, publication status, and priority. Published announcements can notify alumni, and Facebook-connected records show whether the update came from or was synced with the Facebook page. The archive action was replaced with delete so the visible action and database behavior are easier to understand: deleting an announcement also removes its related read records and announcement notifications.",
            "screenshot": "Capture the announcement filter area, announcement table, Facebook synced badge, edit button, and delete button.",
            "connect": "Connects to Objective 3 and Objective 4, especially announcement management and communication support.",
        },
        {
            "title": "5.14 Feedback Management Interface",
            "result": "The admin feedback module organizes alumni feedback records so ARLO can review, filter, update, and monitor submitted concerns.",
            "discussion": "The feedback management interface completes the two-way communication cycle. Alumni submit concerns from the alumni side, and administrators review them from the admin side. This keeps communication traceable and easier to summarize compared with separate messages outside the system. The module supports ARLO's need to document alumni concerns and use those concerns as input for service improvement.",
            "screenshot": "Capture the admin Feedback Management table and one feedback detail/update page.",
            "connect": "Connects to Objective 4 and Functional Requirement: Feedback Management.",
        },
        {
            "title": "5.15 Analytics and Report Interface",
            "result": "The analytics module generates filtered tracer and employment reports using graphs, KPI cards, report previews, and export buttons based on available registered alumni and tracer records.",
            "discussion": "Analytics is a reporting layer, not a separate database table. The system reads existing records and turns them into charts such as employment trends, traced graduates, industry distribution, employability rate, geographic distribution, unemployment reasons, user growth, job submissions, college distribution, employment status, graduation year trends, feedback categories, curriculum relevance, time to first job, and satisfaction indicators. The filters allow administrators to choose date ranges, college, program, sex, civil status, employment status, year graduated, and visible graph sections. If a batch has imported roster records but no registered tracer responses yet, the report properly explains that graphs may be empty because CSV-only records are for verification, not tracer analytics.",
            "screenshot": "Capture the report filters, the report preview header, the registered/not registered explanation, and at least three graph examples with X and Y axis labels.",
            "connect": "Connects to Objective 5 and Technical Background: Chart.js Analytics.",
        },
        {
            "title": "5.16 Audit Logs Interface",
            "result": "The audit logs module records important administrative and account actions to support accountability and traceability.",
            "discussion": "Audit logs help answer the question of who changed what and when. This is important because the system handles alumni personal data, administrator actions, announcements, notifications, job records, account updates, and deletion events. By storing these actions in system_logs, the platform supports the privacy and accountability concerns identified during requirements analysis.",
            "screenshot": "Capture the admin Audit Logs page showing action names, actor, description, date, and source information.",
            "connect": "Connects to Business Rule: Audit Log Monitoring and Non-Functional Requirement: Security.",
        },
        {
            "title": "5.17 Super Admin Account Control Interface",
            "result": "The Super Admin workspace is separated from regular admin operations and focuses on administrator account management and alumni account maintenance.",
            "discussion": "The Super Admin interface is intentionally narrower than the regular admin workspace. It does not manage jobs, announcements, analytics, or feedback as an operational staff member would. Instead, it creates and manages regular administrator accounts and can support account-level maintenance. This matches the updated ERD decision where regular administrators are stored in admins and Super Admin accounts are stored in the separate super_admins table. The separation makes the account governance role clearer and easier to defend in the database design.",
            "screenshot": "Capture the Super Admin admin accounts list, create admin form, edit admin form, and the Super Admin scope cards.",
            "connect": "Connects to Super Admin Role Management Rule and the updated ERD with separate admins and super_admins tables.",
        },
    ]


def add_module_section(doc: Document, block: dict) -> None:
    doc.add_heading(block["title"], level=2)
    add_para(doc, "Result: " + block["result"], bold_label="Result:")
    add_para(doc, "Discussion: " + block["discussion"], bold_label="Discussion:")
    add_para(doc, "Suggested screenshot: " + block["screenshot"], bold_label="Suggested screenshot:")
    add_para(doc, "Connection to the paper: " + block["connect"], bold_label="Connection to the paper:")


def add_screenshot_table(doc: Document) -> None:
    rows = [
        ("Figure 17", "Login and account recovery", "Login page, forgot-password page, professional success or validation message", "Shows secure access and recovery workflow."),
        ("Figure 18", "Registration and email verification", "Alumni verification fields, email code panel, password unlock state", "Shows controlled alumni account creation."),
        ("Figure 19", "Alumni dashboard", "Profile card, tracer status, employment record, checklist, notifications", "Shows alumni home screen and progress guidance."),
        ("Figure 20", "Profile and settings", "Profile update form and password/account settings", "Shows alumni record maintenance."),
        ("Figure 21", "Tracer form", "Employment, salary, industry, proof upload, satisfaction, submit area", "Shows graduate tracer data collection."),
        ("Figure 22", "Jobs and job opportunity notice", "Jobs page, Notify Job Opportunity modal, recent submissions", "Shows that alumni notify opportunities rather than apply."),
        ("Figure 23", "Announcements and notifications", "Notification list, announcement detail, read/unread badge", "Shows communication support."),
        ("Figure 24", "Feedback submission", "Feedback category and message form", "Shows alumni concern submission."),
        ("Figure 25", "Admin dashboard", "Registered and not registered alumni cards, jobs, announcements, this month", "Shows admin overview and count separation."),
        ("Figure 26", "Alumni directory and CSV import", "Directory table, CSV import modal, template comparison", "Shows roster import and account management."),
        ("Figure 27", "Direct alumni notification", "Bell action and Notice/Reminder/Warning modal", "Shows admin-to-alumni notification control."),
        ("Figure 28", "Job management", "Job table, pending submissions, approve/reject/delete controls", "Shows administrator control over job opportunities."),
        ("Figure 29", "Announcement management", "Filters, announcement records, Facebook sync badge, delete action", "Shows announcement administration."),
        ("Figure 30", "Feedback management", "Feedback records and feedback detail/update screen", "Shows admin review of alumni feedback."),
        ("Figure 31", "Analytics and reports", "Filters, report preview, KPI cards, charts, export buttons", "Shows tracer and employment analytics."),
        ("Figure 32", "Audit logs", "System log table with action and actor information", "Shows accountability and traceability."),
        ("Figure 33", "Super Admin workspace", "Admin accounts list, create/edit admin form, scope cards", "Shows separate account governance role."),
    ]
    table = doc.add_table(rows=1, cols=4)
    widths = column_widths_from_weights(9360, [0.9, 1.7, 2.4, 2.0])
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=0)
    headers = ["Figure", "Screen", "What to Capture", "Why It Matters"]
    for i, header in enumerate(headers):
        set_cell_text(table.cell(0, i), header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, bold=(i == 0), color=ACCENT if i == 0 else None, size=9)
    style_table(table, "F4FBF7")


def add_corrections_table(doc: Document) -> None:
    rows = [
        (
            "Job applications still appear in objectives, scope, requirements, DFD, use case, data model, or ERD.",
            "Change to job browsing and alumni job opportunity submission. Remove portal job application wording unless you explicitly restore that feature.",
            "Current live database has no job_applications table.",
        ),
        (
            "The paper says Super Admin and Admin are both stored in admins with a role field.",
            "Use the updated design: admins table for regular administrators and super_admins table for Super Admin accounts.",
            "The current database has a separate super_admins table.",
        ),
        (
            "The paper lists queue_jobs, job_batches, job_matches, or job_applications as active tables.",
            "Remove them from the active database dictionary and ERD. Mention only if you explain them as removed support tables, which is usually unnecessary.",
            "Those tables are not in the current live table list.",
        ),
        (
            "The paper says notifications do not send email automatically.",
            "Update to: notifications are stored in-system and can also be emailed when Gmail/mail settings are configured.",
            "Admin direct notices and announcement notifications use database and mail channels.",
        ),
        (
            "The paper describes announcement archive actions.",
            "Use delete/remove for the current admin announcement UI. Explain that delete removes related read records and announcement notifications.",
            "The current announcement delete action hard-deletes records.",
        ),
        (
            "The paper explains 450 alumni as if all are registered accounts.",
            "Separate roster/imported alumni from registered alumni. Use the wording: CSV-only records support signup verification but do not count as tracer respondents.",
            "Your admin overview separates registered and not registered alumni counts.",
        ),
        (
            "Testing section repeats Table 10 numbering.",
            "Renumber the second Table 10 to Table 11 or continue the correct sequence.",
            "This is a formatting and defense clarity issue.",
        ),
        (
            "Related studies say 'we emphasized' when referring to other authors.",
            "Use 'the authors emphasized' or 'the study emphasized.'",
            "This avoids accidentally claiming another study as your own work.",
        ),
        (
            "References list Laravel Documentation 13.x while the implementation is Laravel 12.48.1.",
            "Cite Laravel 12 documentation or describe it as Laravel documentation used during development.",
            "Your running app reports Laravel 12.48.1.",
        ),
        (
            "Use case still includes Apply for Job.",
            "Replace with Browse Job Opportunities and Submit Job Opportunity for Admin Review.",
            "The alumni-side apply form was removed.",
        ),
    ]
    table = doc.add_table(rows=1, cols=3)
    widths = column_widths_from_weights(9360, [2.3, 2.7, 1.8])
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=0)
    headers = ["What Is Wrong or Outdated", "Correction to Apply", "Reason"]
    for i, header in enumerate(headers):
        set_cell_text(table.cell(0, i), header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, size=9)
        set_cell_shading(cells[0], LIGHT_RED)
    style_table(table, "FFF7F8")


def add_result_summary_table(doc: Document) -> None:
    rows = [
        ("Secure account access", "Login, email verification, password reset, role-specific access", "Objective 1"),
        ("Alumni information management", "Profile, settings, tracer form, employment record, scholarship/further study fields", "Objective 2"),
        ("Communication and engagement", "Announcements, notifications, direct admin notices, feedback, job opportunity notices", "Objective 3"),
        ("Administrative control", "Admin dashboard, alumni directory, CSV import, jobs, announcements, feedback, audit logs", "Objective 4"),
        ("Analytics and reporting", "Filters, KPI cards, charts, report preview, PDF/CSV export", "Objective 5"),
        ("Usability support", "Guided dashboards, clear cards, modals, validation messages, status separation", "SUS objective"),
    ]
    table = doc.add_table(rows=1, cols=3)
    widths = column_widths_from_weights(9360, [1.8, 3.8, 1.2])
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=0)
    for i, header in enumerate(["Result Area", "Implemented UI / Feature Output", "Paper Link"]):
        set_cell_text(table.cell(0, i), header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, bold=(i == 0), color=ACCENT if i == 0 else None)
    style_table(table, "F4FBF7")


def build() -> None:
    doc = Document()
    setup_styles(doc)

    p = doc.add_paragraph(style="Title")
    p.add_run("CHAPTER V RESULTS AND DISCUSSION").bold = True
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("UI, Features, and Modules Writing Guide for the CMU OARL Web-Based Tracer and Employment Analytics System")
    add_metadata_table(doc)
    add_callout(
        doc,
        "How to use this file",
        "Use this document as your Chapter V guide. Insert your own screenshots manually, then keep the Result and Discussion text beside the correct figure. The wording is aligned with the current system: separate admins and super_admins tables, no job application table, no queue_jobs/job_batches tables, and alumni job handling through job opportunity notices.",
        LIGHT_BLUE,
    )

    doc.add_page_break()
    add_header_footer(doc)

    doc.add_heading("1. Current System Verification", level=1)
    add_para(
        doc,
        "Before writing the results, the current implementation was checked against the Laravel routes, views, controllers, notifications, and PostgreSQL schema. This matters because some older parts of the manuscript still describe features that have already been changed in the system."
    )
    add_system_verification_table(doc)

    doc.add_heading("2. Overall Result", level=1)
    add_para(
        doc,
        "Result: The developed system produced a centralized web-based platform for alumni verification, registered alumni account management, tracer data submission, employment opportunity sharing, announcements, notifications, feedback, analytics, audit logs, administrator operations, and Super Admin account control.",
        bold_label="Result:",
    )
    add_para(
        doc,
        "Discussion: The system addressed the main problem of scattered and manual alumni tracer management by moving the workflow into one Laravel and PostgreSQL platform. Alumni no longer depend only on open Google Forms to submit tracer information, and administrators no longer depend only on spreadsheets to organize graduate records. The interface separates alumni actions from administrative actions, and it also separates regular administrator accounts from Super Admin accounts. This makes the system easier to explain because each actor has a clear responsibility: alumni maintain and submit their own information, administrators manage operational records and reports, and the Super Admin governs administrator access.",
        bold_label="Discussion:",
    )
    add_result_summary_table(doc)

    doc.add_heading("3. Results and Discussion by UI Module", level=1)
    add_para(
        doc,
        "The following sections are written in a result-first format. Each module begins with what the system produced, followed by a discussion of what the user sees, how the feature works, and why it matters to the study."
    )
    for block in module_blocks():
        add_module_section(doc, block)

    doc.add_page_break()
    doc.add_heading("4. Manual Screenshot Guide", level=1)
    add_para(
        doc,
        "Use this checklist when collecting screenshots for Chapter V. Capture the screens in the same order as the discussion so the figures feel like a walkthrough of the actual system."
    )
    add_screenshot_table(doc)

    doc.add_heading("5. Corrections Needed in the Current Paper", level=1)
    add_callout(
        doc,
        "Important consistency note",
        "Your paper is strong in structure, but several parts still describe older features. The biggest correction is the job module: the current system no longer has alumni job applications. It has job browsing and alumni-submitted job opportunities for administrator review.",
        LIGHT_AMBER,
    )
    add_corrections_table(doc)

    doc.add_heading("6. Suggested Chapter V Opening Paragraph", level=1)
    add_para(
        doc,
        "The implementation of the Web-Based Tracer and Employment Analytics System for the Alumni Relations and Linkages Office resulted in a centralized platform that supports alumni verification, profile management, tracer submission, job opportunity information, announcements, notifications, feedback, analytics, reports, audit logs, and administrator account control. The system interface was organized according to the roles of Alumni, Administrator, and Super Admin so that each user group could access only the functions relevant to its responsibility. The following sections discuss the major user interfaces, implemented modules, and system outputs based on the objectives and requirements of the study."
    )

    doc.add_heading("7. Suggested Closing Paragraph for Results and Discussion", level=1)
    add_para(
        doc,
        "Overall, the implemented system showed that the manual tracer workflow of the Alumni Relations and Linkages Office could be transformed into an organized web-based environment. The alumni interface supported registration, verification, profile maintenance, tracer submission, announcements, notifications, job opportunity viewing, job opportunity submission, and feedback. The administrator interface supported alumni record management, CSV roster import, job posting management, announcement management, feedback handling, analytics, reporting, and audit logging. The Super Admin interface provided a restricted workspace for administrator account control. These results show that the developed platform met the functional objectives of the study while also improving record organization, communication, reporting readiness, and accountability."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
